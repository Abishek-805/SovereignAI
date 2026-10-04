from pathlib import Path
import pytest
from openpyxl import load_workbook
from docx import Document
from workflows.generated_documents import generate
from workflows.project_map import inspect_project
from router.task_completion import observe_completion
from backend.contracts import WorkbenchError


def test_workbook_preserves_formula_and_offline_value(tmp_path):
    path=tmp_path/'budget.xlsx'
    checks=generate(path,{'sheets':[{'name':'Budget','rows':[['Item','Cost'],['A',10],['B',20],['Total','=SUM(B2:B3)']]}]})
    assert checks['computed_formula_values']=={'Budget!B4':30}
    assert load_workbook(path,data_only=False)['Budget']['B4'].value=='=SUM(B2:B3)'
    assert load_workbook(path,data_only=True)['Budget']['B4'].value==30
    with pytest.raises(WorkbenchError):generate(path,{'sheets':[]})
    assert load_workbook(path,data_only=True)['Budget']['B4'].value==30


def test_invalid_reference_is_rejected_with_cell_evidence(tmp_path):
    path=tmp_path/'invalid.xlsx'
    with pytest.raises(WorkbenchError,match='B2=10'):
        generate(path,{'sheets':[{'name':'Budget','rows':[['Item','Cost'],['A',10],['Total','=SUM(C2:C3)']]}]})
    assert not path.exists()


def test_docx_content_and_table_readback(tmp_path):
    path=tmp_path/'guide.docx'
    generate(path,{'title':'Guide','sections':[{'heading':'Checks','paragraphs':['Real results'], 'tables':[[['Step','Result'],['Run','PASS']]]}]})
    doc=Document(path)
    assert any(p.text=='Real results' for p in doc.paragraphs)
    assert doc.tables[0].cell(1,1).text=='PASS'


def test_project_map_uses_unique_nested_package():
    result=inspect_project({'web/package.json':b'{"scripts":{"build":"vite build"},"devDependencies":{"vite":"1"}}','web/index.html':b''})
    assert result['cwd']=='web'
    assert result['commands']['run']=='npm run build --offline'
    assert result['framework']=='Vite'
    assert not inspect_project({'main.py':b'', 'app.py':b''})['commands'].get('run')


def test_dependency_block_is_not_model_failure():
    result=observe_completion({'state':'blocked','exit_code':127,'blocked_reason':'DEPENDENCY','blocker':'vite: not found'})
    assert result['state']=='blocked'
    assert not result['achieved']
