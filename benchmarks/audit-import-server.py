"""Disposable browser acceptance server: never uses canonical user data."""
import tempfile
import time
import sys
from pathlib import Path
from types import SimpleNamespace
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import uvicorn
from starlette.responses import JSONResponse
from starlette.responses import Response
from io import BytesIO
from backend.app import create_app
from backend.service import Workbench
from backend.settings import Settings
from rag.store import Store
from tests.test_service import TestEmbedder
from tests.test_answer import FakeModel
from router.sandbox import SandboxResult

class FixtureSandbox:
    """UI contract fixture only; this does not verify Docker isolation."""
    image_id='disposable-test-fixture'
    def execute(self,code,input_files):
        return SandboxResult(0,'Fixture validation passed','',True)

class SlowEmbedder(TestEmbedder):
    def encode(self, texts, query=False):
        time.sleep(.5)
        return super().encode(texts, query)

with tempfile.TemporaryDirectory(prefix='sovereign-import-browser-') as directory:
    root=Path(directory)
    model=FakeModel(); model.status=lambda: {'available': True, 'is_sleeping': False}
    model.plan_task=lambda *args,**kwargs:{'action':'application_tools','target':'','expression':'','response':'','document_scope':'focused'}
    model.plan_application_tools=lambda *args,**kwargs:[{'tool':'file_delete_scope','target':'all','value':'','input':''}]
    service=Workbench(Settings(data_dir=root/'data', project_dir=root/'projects'),
                      Store(root/'index.sqlite'), SlowEmbedder(), model,
                      SimpleNamespace(acquire_lease=lambda capability: None))
    service._verified_coding_sandbox=lambda: FixtureSandbox()
    app=create_app(service)
    server=None
    @app.middleware('http')
    async def fixture_identity(request,call_next):
        if request.url.path=='/__audit_workbook__':
            from openpyxl import Workbook
            from openpyxl.styles import Font,PatternFill
            book=Workbook();sheet=book.active;sheet.title='Results'
            sheet.append(['Example results']);sheet.merge_cells('A1:C1');sheet.append(['Name','Score','Result'])
            sheet['A2'].font=Font(bold=True);sheet['A2'].fill=PatternFill('solid',fgColor='FFCCEECC')
            for i in range(150):sheet.append([f'Learner {i}',i,'PASS'])
            book.create_sheet('Notes').append(['Original note'])
            output=BytesIO();book.save(output);response=Response(output.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        elif request.url.path=='/__audit_shutdown__' and request.method=='POST':
            server.should_exit=True
            response=JSONResponse({'stopping':True})
        else:
            response=await call_next(request)
        response.headers['x-sovereign-audit']='disposable'
        return response
    server=uvicorn.Server(uvicorn.Config(app,host='127.0.0.1',port=8088,log_level='error'))
    server.run()
