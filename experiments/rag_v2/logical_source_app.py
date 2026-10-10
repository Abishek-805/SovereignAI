"""Isolated actual Agent acceptance with independently authored workbook values."""
from pathlib import Path
import os,json,time,hashlib,sys,importlib.util
import psutil

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'evidence/logical-source-resolution/20261009'


def app():
    arm=os.environ.get('SOVEREIGN_LOGICAL_ARM','after')
    names=['rag/tables.py','rag/identifiers.py','rag/analytics.py','backend/service.py']
    if arm=='before':
        for name in names:
            key=name[:-3].replace('/','.')
            spec=importlib.util.spec_from_file_location(key,OUT/'rollback'/name)
            module=importlib.util.module_from_spec(spec);sys.modules[key]=module;spec.loader.exec_module(module)
    from backend.app import create_app
    from backend.service import Workbench
    from backend.settings import Settings
    from backend.model import LocalModel
    from openpyxl import Workbook
    from rag.tables import _load
    directory=OUT/arm;directory.mkdir(exist_ok=True)
    fixtures=OUT/'fixtures';fixtures.mkdir(exist_ok=True)
    main=fixtures/'Synthetic Professional Skills Training All Marks.xlsx'
    if not main.exists():
        book=Workbook();book.active.title='All'
        for row in [['Member Id','WAT 1','WAT 2','CAT 1','WAT 3'],['24ALR056',7,6,12.5,'AB'],['24ALR001',14,20,30,'AB']]:book.active.append(row)
        for name,values in [('Week 2',[6,20]),('Week 3',[7,14])]:
            sheet=book.create_sheet(name);sheet.append(['Member Id','Score','Result'])
            for entity,value in zip(['24ALR056','24ALR001'],values):sheet.append([entity,value,'Recorded'])
        book.save(main)
    other=fixtures/'Synthetic Service Results.csv'
    if not other.exists():other.write_text('Member Id,WAT 1,WAT 2\n24ALR001,99,88\n')
    settings=Settings(data_dir=directory/'ui-data',context=16384,output_tokens=1024)
    model=LocalModel(verify_semantics=False);wb=Workbench(settings,model=model)
    for path in [main,other]:wb.import_file(path,display_name=path.name)
    _load.cache_clear();run=str(time.time_ns());snapshot=directory/run;snapshot.mkdir()
    for name in names+(['rag/table_resolution.py'] if arm=='after' else []):
        source=OUT/'rollback'/name if arm=='before' else ROOT/name
        target=snapshot/name;target.parent.mkdir(parents=True,exist_ok=True);target.write_bytes(source.read_bytes())
    (snapshot/'runtime.json').write_text(json.dumps({'arm':arm,'run_id':run,'python':sys.version,'model_server':'existing 8087',
        'context':16384,'output_tokens':1024,'source_hashes':{str(p.relative_to(snapshot)):hashlib.sha256(p.read_bytes()).hexdigest() for p in snapshot.rglob('*.py')}},indent=2))
    calls=[];request=model._request
    def observed(method,path,**kwargs):
        start=time.perf_counter()
        try:return request(method,path,**kwargs)
        finally:calls.append({'method':method,'path':path,'seconds':time.perf_counter()-start})
    model._request=observed
    execute=wb.run_auto_agent
    def agent(goal,*args,**kwargs):
        start=time.perf_counter();offset=len(calls);process=psutil.Process()
        receipt={'run_id':run,'arm':arm,'question':goal,'rss_before':process.memory_info().rss,
            'cache_before':_load.cache_info()._asdict(),'history':args[2] if len(args)>2 else kwargs.get('history'),
            'evidence_ids':kwargs.get('evidence_ids')}
        try:
            result=execute(goal,*args,**kwargs);receipt['result']=result;return result
        except Exception as exc:
            receipt['error']=str(exc);receipt['code']=getattr(exc,'code',None);raise
        finally:
            receipt.update(seconds=time.perf_counter()-start,rss_after=process.memory_info().rss,cache_after=_load.cache_info()._asdict(),model_requests=calls[offset:])
            with (directory/'trajectories.jsonl').open('a',encoding='utf-8') as handle:handle.write(json.dumps(receipt,ensure_ascii=False)+'\n')
    wb.run_auto_agent=agent
    result=create_app(wb)
    staging=ROOT/'evidence/universal-report/20261009/frontend-dist'
    from fastapi.staticfiles import StaticFiles
    result.router.routes[-1].app=StaticFiles(directory=str(staging),html=True)
    return result
