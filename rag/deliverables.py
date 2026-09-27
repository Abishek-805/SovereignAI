import json
import logging
import os
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import re
from tempfile import TemporaryDirectory

from backend.contracts import WorkbenchError
from backend.settings import Settings

logger = logging.getLogger(__name__)

def _sanitize_filename(name: str) -> str:
    import re
    # Remove any directory traversals or invalid characters
    name = re.sub(r'[^\w\-. ]', '_', name)
    name = name.strip()
    if not name:
        name = "document"
    return name

def _sanitize_excel_cell(value) -> str:
    if isinstance(value, str):
        val = value.strip()
        if val.startswith(('=', '+', '-', '@')):
            return f"'{val}"
    return value

def validate_deliverable_schema(data: dict):
    required = ['task_id', 'title', 'document_date', 'source_versions', 
                'findings', 'calculations', 'requested_action', 
                'open_questions', 'references']
    for req in required:
        if req not in data:
            raise WorkbenchError('invalid_schema', f"Missing required field: {req}")

def write_deliverables(task_id: str, data: dict, settings: Settings) -> dict:
    validate_deliverable_schema(data)
    if not isinstance(task_id,str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,80}',task_id) or data['task_id']!=task_id:
        raise WorkbenchError('invalid_task','Task ID must be a matching safe identifier')
    output_root = settings.data_dir.parent / 'outputs'
    output_root.mkdir(parents=True,exist_ok=True)
    outputs_dir=output_root/task_id
    if outputs_dir.exists():
        raise WorkbenchError('artifact_exists','Task outputs already exist')
    
    word_filename = f"{_sanitize_filename(data['title'])}.docx"
    excel_filename = f"{_sanitize_filename(data['title'])}_evidence.xlsx"
    slide_filename = f"{_sanitize_filename(data['title'])}_briefing.pptx"
    
    try:
        with TemporaryDirectory(prefix='.building-',dir=output_root) as temporary:
            temp=Path(temporary)
            word_path=temp/word_filename
            excel_path=temp/excel_filename
            slide_path=temp/slide_filename
            generate_word_draft(data, word_path)
            generate_excel_evidence(data, excel_path)
            generate_slide_briefing(data, slide_path)
            from docx import Document
            from openpyxl import load_workbook
            from pptx import Presentation
            document=Document(word_path)
            workbook=load_workbook(excel_path,read_only=True,data_only=True)
            try:
                slides=Presentation(slide_path)
                slide_text=' '.join(shape.text for slide in slides.slides for shape in slide.shapes if shape.has_text_frame)
                if (not document.paragraphs or not {'Findings','Calculations','Sources','Unresolved Issues'} <= set(workbook.sheetnames)
                        or len(slides.slides)<3 or not all(ref['id'] in slide_text for ref in data['references'])):
                    raise WorkbenchError('artifact_invalid','Generated files are missing required sections')
            finally:
                workbook.close()
            metadata={'task_id':task_id,'created_at':datetime.now(timezone.utc).isoformat(),
                      'source_versions':data['source_versions'],'checks':{'docx_reopened':True,'xlsx_reopened':True,'pptx_reopened':True},
                      'files':[{ 'name':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                                'bytes':path.stat().st_size} for path in (word_path,excel_path,slide_path)]}
            (temp/'manifest.json').write_text(json.dumps(metadata,indent=2),encoding='utf-8')
            temp.rename(outputs_dir)
    except WorkbenchError:
        raise
    except Exception as e:
        logger.error(f"Failed to generate deliverables: {e}")
        raise WorkbenchError('generation_failed', f"Deliverable generation failed: {e}")
        
    return {
        "task_id": task_id,
        "word_path": str((outputs_dir/word_filename).resolve()),
        "excel_path": str((outputs_dir/excel_filename).resolve()),
        "slide_path": str((outputs_dir/slide_filename).resolve())
    }


def generate_slide_briefing(data: dict, out_path: Path):
    """Create a compact evidence-linked briefing from checked structured findings."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation=Presentation()
    presentation.slide_width=Inches(13.333)
    presentation.slide_height=Inches(7.5)

    def slide(title, lines):
        page=presentation.slides.add_slide(presentation.slide_layouts[6])
        header=page.shapes.add_textbox(Inches(.7), Inches(.45), Inches(12), Inches(.7))
        header.text_frame.paragraphs[0].text=title
        header.text_frame.paragraphs[0].font.size=Pt(26)
        body=page.shapes.add_textbox(Inches(.8), Inches(1.35), Inches(11.8), Inches(5.75))
        frame=body.text_frame
        frame.word_wrap=True
        for index,line in enumerate(lines):
            paragraph=frame.paragraphs[0] if index==0 else frame.add_paragraph()
            paragraph.text=str(line)[:950]
            paragraph.font.size=Pt(16)
            paragraph.space_after=Pt(16)

    slide(data['title'], [f"Draft briefing · {data['document_date']}",
                          f"Task {data['task_id']}", data['requested_action']])
    findings=data['findings'] or []
    for start in range(0,len(findings),2):
        lines=[]
        for finding in findings[start:start+2]:
            lines.extend([f"Observation: {finding.get('observation','')}",
                          f"Requirement: {finding.get('requirement','')}",
                          f"Status: {finding.get('status','')} · Sources: {', '.join(finding.get('source_ids',[]))}"])
        slide(f"Findings {start+1}–{min(start+2,len(findings))}",lines)
    issues=data['open_questions'] or ['None recorded.']
    for start in range(0,len(issues),5):
        slide('Unresolved issues',issues[start:start+5])
    references=[f"[{ref['id']}] {ref['description']}" for ref in data['references']]
    for start in range(0,len(references),5):
        slide('Source references',references[start:start+5])
    presentation.save(out_path)

def generate_word_draft(data: dict, out_path: Path):
    try:
        import docx
        from docx.shared import Pt, Inches
    except ImportError as e:
        raise WorkbenchError('renderer_unavailable', f"python-docx is not available: {e}")

    doc = docx.Document()
    from docx.enum.section import WD_ORIENT
    section=doc.sections[0]
    section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width,section.page_height=section.page_height,section.page_width
    section.top_margin=section.bottom_margin=Inches(0.65)
    section.left_margin=section.right_margin=Inches(0.7)
    
    # Title
    doc.add_heading(data['title'], 0)
    
    # Metadata
    doc.add_paragraph(f"Date: {data['document_date']}")
    doc.add_paragraph(f"Task ID: {data['task_id']}")
    
    # Executive Summary (using requested action)
    doc.add_heading("Requested Action", level=1)
    doc.add_paragraph(data['requested_action'])
    
    # Findings Table
    doc.add_heading("Findings", level=1)
    if data['findings']:
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        table.autofit=False
        for row in table.rows:
            for cell,width in zip(row.cells,(3.15,3.0,1.75,0.75)):
                cell.width=Inches(width)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Observation'
        hdr_cells[1].text = 'Requirement'
        hdr_cells[2].text = 'Status'
        hdr_cells[3].text = 'Sources'
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        
        for finding in data['findings']:
            row_cells = table.add_row().cells
            row_cells[0].text = str(finding.get('observation', ''))
            row_cells[1].text = str(finding.get('requirement', ''))
            row_cells[2].text = str(finding.get('status', ''))
            row_cells[3].text = ", ".join(finding.get('source_ids', []))
            for cell,width in zip(row_cells,(3.15,3.0,1.75,0.75)):
                cell.width=Inches(width)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size=Pt(9)
    else:
        doc.add_paragraph("No findings.")
        
    # Open Questions
    doc.add_heading("Unresolved Issues", level=1)
    if data['open_questions']:
        for q in data['open_questions']:
            doc.add_paragraph(q, style='List Bullet')
    else:
        doc.add_paragraph("None.")
        
    # References
    doc.add_heading("References", level=1)
    for ref in data['references']:
        doc.add_paragraph(f"[{ref.get('id', '')}] {ref.get('description', '')}")
        
    doc.save(out_path)

def generate_excel_evidence(data: dict, out_path: Path):
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise WorkbenchError('renderer_unavailable', f"openpyxl is not available: {e}")
        
    wb = Workbook()
    
    # Sheet 1: Findings
    ws_findings = wb.active
    ws_findings.title = "Findings"
    ws_findings.append(["Observation", "Requirement", "Status", "Source IDs", "Calc ID"])
    for f in data.get('findings', []):
        ws_findings.append([
            _sanitize_excel_cell(f.get('observation')),
            _sanitize_excel_cell(f.get('requirement')),
            _sanitize_excel_cell(f.get('status')),
            _sanitize_excel_cell(", ".join(f.get('source_ids', []))),
            _sanitize_excel_cell(f.get('calc_id'))
        ])
        
    # Sheet 2: Calculations
    ws_calcs = wb.create_sheet("Calculations")
    ws_calcs.append(["Calc ID", "Expression", "Result", "Rounded"])
    for c in data.get('calculations', []):
        ws_calcs.append([
            _sanitize_excel_cell(c.get('id')),
            _sanitize_excel_cell(c.get('expression')),
            c.get('result'),
            c.get('rounded')
        ])
        
    # Sheet 3: Sources
    ws_src = wb.create_sheet("Sources")
    ws_src.append(["Document ID", "Version Hash"])
    for s in data.get('source_versions', []):
        ws_src.append([
            _sanitize_excel_cell(s.get('document_id')),
            _sanitize_excel_cell(s.get('version_hash'))
        ])
        
    # Sheet 4: Unresolved Issues
    ws_issues = wb.create_sheet("Unresolved Issues")
    ws_issues.append(["Issue"])
    for i in data.get('open_questions', []):
        ws_issues.append([_sanitize_excel_cell(i)])

    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    widths={'Findings':[52,52,34,16,14], 'Calculations':[16,26,16,16],
            'Sources':[68,68], 'Unresolved Issues':[90]}
    for sheet in wb.worksheets:
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for index,width in enumerate(widths[sheet.title],1):
            sheet.column_dimensions[get_column_letter(index)].width=width
        for cell in sheet[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='283A32')
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
            sheet.row_dimensions[row[0].row].height=42 if sheet.title in {'Findings','Unresolved Issues'} else 30
        
    wb.save(out_path)
