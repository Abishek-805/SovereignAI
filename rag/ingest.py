import hashlib
import statistics
from io import BytesIO
from pathlib import Path
from bisect import bisect_left
import posixpath
from zipfile import ZipFile, BadZipFile
from lxml import etree
from pypdf import PdfReader
from backend.contracts import WorkbenchError, Page, Chunk, Extraction
from rag.pdf_visuals import overall_star_rating

TEXT_FORMATS = {'.txt', '.md', '.markdown', '.csv', '.tsv', '.json', '.jsonl', '.log',
                '.xml', '.html', '.htm', '.css', '.yaml', '.yml', '.toml', '.ini', '.cfg',
                '.sql', '.py', '.js', '.jsx', '.ts', '.tsx', '.java', '.c', '.cpp', '.h',
                '.cs', '.go', '.rs', '.php', '.rb', '.sh', '.ps1', '.tex'}
SUPPORTED = TEXT_FORMATS | {'.pdf', '.docx', '.xlsx', '.pptx'}


def _xlsx_pages(data: bytes, max_sheets: int):
    """Read actual stored cells, never expand a worksheet's declared dimensions."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
    with ZipFile(BytesIO(data)) as archive:
        def xml(name):
            return etree.fromstring(archive.read(name), parser=parser)
        shared = []
        if 'xl/sharedStrings.xml' in archive.namelist():
            shared = [''.join(item.xpath('.//s:t/text()', namespaces=ns)) for item in xml('xl/sharedStrings.xml')]
        relations = {item.get('Id'): item.get('Target') for item in xml('xl/_rels/workbook.xml.rels')
                     if item.get('TargetMode') != 'External'}
        sheets = xml('xl/workbook.xml').findall('s:sheets/s:sheet', ns)
        if len(sheets) > max_sheets:
            raise WorkbenchError('page_limit', 'Spreadsheet exceeds sheet limit')
        pages, count, total = [], 0, 0
        for sheet in sheets:
            rel = sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            target = relations[rel]
            name = posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/' + target)
            if not name.startswith('xl/') or '..' in name.split('/'):
                raise WorkbenchError('parse_failed', 'Invalid spreadsheet worksheet path')
            lines = []
            with archive.open(name) as source:
                for _, cell in etree.iterparse(source, events=('end',), tag='{'+ns['s']+'}c',
                                               resolve_entities=False, no_network=True, huge_tree=False):
                    count += 1
                    if count > 100000:
                        raise WorkbenchError('file_too_large', 'Spreadsheet exceeds 100,000 stored cells')
                    value = cell.findtext('s:v', namespaces=ns)
                    kind = cell.get('t')
                    if kind == 's' and value is not None:
                        index = int(value)
                        if not 0 <= index < len(shared):
                            raise WorkbenchError('parse_failed', 'Invalid spreadsheet shared-string reference')
                        value = shared[index]
                    elif kind == 'inlineStr':
                        value = ''.join(cell.xpath('.//s:t/text()', namespaces=ns))
                    if value is not None:
                        line = f'{cell.get("r", "cell")}: {value}'
                        total += len(line) + 1
                        if total > 2_000_000:
                            raise WorkbenchError('file_too_large', 'Extracted text exceeds two million characters')
                        lines.append(line)
                    cell.clear()
                    while cell.getprevious() is not None:
                        del cell.getparent()[0]
            if lines:
                pages.append(Page('Sheet: '+sheet.get('name', '')+'\n'+'\n'.join(lines), None, method='spreadsheet_cells'))
        return pages


def _docx_text(data: bytes) -> str:
    """Read body paragraphs, table cells, and visible equation text in order."""
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries=archive.infolist()
            if len(entries)>2000 or sum(item.file_size for item in entries)>80*1024*1024:
                raise WorkbenchError('file_too_large','DOCX archive expands beyond the allowed limit')
            document=archive.getinfo('word/document.xml')
            if document.file_size>16*1024*1024:
                raise WorkbenchError('file_too_large','DOCX document XML exceeds 16 MB')
            xml=archive.read(document)
        root=etree.fromstring(xml,parser=etree.XMLParser(resolve_entities=False,no_network=True,huge_tree=False))
    except WorkbenchError:
        raise
    except (BadZipFile,KeyError,OSError,ValueError,RuntimeError,etree.XMLSyntaxError) as exc:
        raise WorkbenchError('parse_failed','DOCX extraction failed') from exc
    word='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    math='http://schemas.openxmlformats.org/officeDocument/2006/math'
    lines=[]
    for paragraph in root.xpath('.//w:body//w:p',namespaces={'w':word}):
        parts=[]
        for node in paragraph.iter():
            if node.tag in {f'{{{word}}}t',f'{{{math}}}t'}:
                parts.append(node.text or '')
            elif node.tag==f'{{{word}}}tab':
                parts.append('\t')
            elif node.tag==f'{{{word}}}br':
                parts.append('\n')
        line=''.join(parts).strip()
        if line: lines.append(line)
    return '\n'.join(lines)


def extract(path: Path, max_bytes=20*1024*1024, max_pages=200) -> Extraction:
    path = Path(path)
    if path.suffix.lower() not in SUPPORTED:
        raise WorkbenchError('unsupported_file', 'Use PDF, DOCX, XLSX, PPTX, CSV, JSON, LOG, UTF-8 TXT or Markdown')
    try:
        with path.open('rb') as source:
            data = source.read(max_bytes + 1)
    except OSError as exc:
        raise WorkbenchError('parse_failed', 'Could not read the selected file') from exc
    if len(data) > max_bytes:
        raise WorkbenchError('file_too_large', 'File exceeds the configured byte limit')
    digest = hashlib.sha256(data).hexdigest()
    pages, warnings = [], []
    if path.suffix.lower() == '.pdf':
        try:
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                raise WorkbenchError('encrypted_pdf', 'Import an unencrypted copy of the PDF')
            if len(reader.pages) > max_pages:
                raise WorkbenchError('page_limit', 'PDF exceeds the configured page limit')
            total = 0
            for number, page in enumerate(reader.pages, 1):
                text = (page.extract_text() or '').replace('\r\n', '\n').replace('\r', '\n')
                if 'Overall Rating' in text:
                    rating = overall_star_rating(data, number)
                    if rating:
                        text += f'\n[Visual PDF evidence: {rating}]'
                total += len(text)
                if total > 2_000_000:
                    raise WorkbenchError('file_too_large', 'Extracted text exceeds two million characters')
                if text.strip() and (len(text.strip()) >= 40 or not page.images):
                    pages.append(Page(text, number))
                else:
                    try:
                        from rag.ocr import render_pdf_page, ocr_image, extract_text_from_observations
                        img = render_pdf_page(data, number, dpi=150)
                        obs = ocr_image(img)
                        ocr_text = extract_text_from_observations(obs)
                        if ocr_text.strip():
                            total += len(ocr_text)
                            if total > 2_000_000:
                                raise WorkbenchError('file_too_large', 'Extracted text exceeds two million characters')
                            confidence=statistics.mean(o.confidence for o in obs)
                            image_hash=hashlib.sha256(img.tobytes()).hexdigest()
                            pages.append(Page(ocr_text, number, method='ocr',confidence=confidence,image_hash=image_hash,
                                              observations=tuple({'text':o.text,'confidence':o.confidence,'box':o.box,'engine':o.engine} for o in obs)))
                            if confidence < 80:
                                warnings.append(f'Page {number}: low OCR confidence; verify critical fields against the image')
                        else:
                            warnings.append(f'Page {number}: no usable text layer; OCR or source inspection needed')
                    except WorkbenchError:
                        raise
                    except Exception as e:
                        warnings.append(f'Page {number}: OCR rendering failed ({str(e)})')
                    if text.strip() and not any(item.page == number for item in pages):
                        pages.append(Page(text, number))
                        warnings.append(f'Page {number}: only a sparse text layer could be read; inspect the original')
        except WorkbenchError:
            raise
        except Exception as exc:
            raise WorkbenchError('parse_failed', 'PDF text extraction failed') from exc
    elif path.suffix.lower() in {'.xlsx', '.pptx'}:
        try:
            with ZipFile(BytesIO(data)) as archive:
                if len(archive.infolist()) > 4000 or sum(item.file_size for item in archive.infolist()) > 80*1024*1024:
                    raise WorkbenchError('file_too_large','Office archive expands beyond 80 MB')
            if path.suffix.lower()=='.xlsx':
                pages.extend(_xlsx_pages(data, max_pages))
                warnings.append('Spreadsheet formulas use saved values; numeric dates and numbers retain their stored values. This import does not recalculate formulas or interpret charts.')
            else:
                from pptx import Presentation
                presentation=Presentation(BytesIO(data))
                if len(presentation.slides)>max_pages: raise WorkbenchError('page_limit','Presentation exceeds slide limit')
                for index,slide in enumerate(presentation.slides,1):
                    lines=[shape.text for shape in slide.shapes if shape.has_text_frame]
                    for shape in slide.shapes:
                        if shape.has_table: lines.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
                    if any(lines): pages.append(Page('\n'.join(lines),index,method='slide_text'))
                warnings.append('Slide text and tables imported; embedded images are not interpreted automatically.')
            if sum(len(page.text) for page in pages)>2_000_000: raise WorkbenchError('file_too_large','Extracted text exceeds two million characters')
        except WorkbenchError: raise
        except Exception as exc: raise WorkbenchError('parse_failed','Could not read this Office file') from exc
    elif path.suffix.lower()=='.docx':
        text=_docx_text(data)
        if len(text)>2_000_000:
            raise WorkbenchError('file_too_large','Extracted text exceeds two million characters')
        if text.strip(): pages.append(Page(text,None,method='docx_text'))
    else:
        try:
            encoding = 'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
            text = data.decode(encoding).replace('\r\n', '\n').replace('\r', '\n')
        except UnicodeDecodeError as exc:
            raise WorkbenchError('invalid_encoding', 'Save the text file as UTF-8') from exc
        if '\x00' in text:
            raise WorkbenchError('invalid_encoding', 'This file contains binary data; import a text export instead')
        if len(text) > 2_000_000:
            raise WorkbenchError('file_too_large', 'Text exceeds two million characters')
        if text.strip():
            pages.append(Page(text, None, method='utf16_text' if encoding == 'utf-16' else 'utf8_text'))
    if not pages:
        raise WorkbenchError('no_extractable_text', 'No usable text found; scanned PDFs need OCR')
    return Extraction(pages, warnings, digest)


def chunk_pages(extraction, tokenizer, document_id, display_name, size=384, overlap=48):
    if size <= overlap or overlap < 0:
        raise ValueError('Chunk size must exceed nonnegative overlap')
    chunks = []
    for page in extraction.pages:
        newlines = [index for index, char in enumerate(page.text) if char == '\n']
        offsets = [(a,b) for a,b in tokenizer.encode(page.text, add_special_tokens=False).offsets if b>a]
        start = 0
        while start < len(offsets):
            end = min(start + size, len(offsets))
            a, b = offsets[start][0], offsets[end-1][1]
            text = page.text[a:b]
            first = page.line_base + bisect_left(newlines, a) if page.page is None else None
            last = first + text.rstrip('\n').count('\n') if first is not None else None
            identity = f'{document_id}:{extraction.source_hash}:{page.page}:{first}:{a}:{b}'
            chunks.append(Chunk(hashlib.sha256(identity.encode()).hexdigest(), document_id,
                                extraction.source_hash, display_name, text, page.page, first, last,
                                page.method,page.confidence,page.image_hash))
            if end == len(offsets):
                break
            start += size - overlap
    if not chunks:
        raise WorkbenchError('no_extractable_text', 'No searchable tokens in source')
    return chunks
