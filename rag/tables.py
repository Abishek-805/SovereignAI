"""Original table previews and bounded, read-only queries over complete records.

The model selects a table/columns/operators; it never supplies executable code.
Source values remain unchanged. Presentation and retrieval use the same schema.
"""
from pathlib import Path
from functools import lru_cache
from zipfile import ZipFile
import csv
import json
import re
import math
import unicodedata
from collections import Counter
from backend.contracts import WorkbenchError

TABULAR={'.xlsx','.csv','.tsv','.json','.jsonl'}

def numeric_value(value):
    """Finite numeric measurement; source cells remain untouched."""
    if isinstance(value,bool):return None
    try:number=float(value)
    except (TypeError,ValueError,OverflowError):return None
    if not math.isfinite(number):return None
    return number


def identity_value(value):
    return unicodedata.normalize('NFC',str(value)).strip().casefold()


def exact_identity_columns(table, entities):
    """Discover complete values before considering a proposed field name."""
    discovered={}
    for column in table['columns']:
        counts=Counter(identity_value(row.get(column)) for row in table['records'])
        matches={str(entity):counts[identity_value(entity)] for entity in entities}
        if any(matches.values()):discovered[column]=matches
    return discovered


def resolve_query_table(plan,available,question):
    identifiers=[token for token in re.findall(r'(?<!\w)[\w]+(?:[.-][\w]+)*(?!\w)',question)
                 if any(c.isalpha() for c in token) and any(c.isdigit() for c in token)]
    if not identifiers or plan.get('table') not in available:return plan
    if exact_identity_columns(available[plan['table']][1],identifiers):return plan
    candidates=[key for key,(_,table) in available.items() if exact_identity_columns(table,identifiers) and
                all(column in table['columns'] for column in plan.get('columns',[]))]
    if len(candidates)>1:raise WorkbenchError('needs_input','Exact identifiers exist in multiple source tables; specify the assessment or source.')
    return {**plan,'table':candidates[0]} if candidates else plan

def scalar(value):
    if isinstance(value,float) and not math.isfinite(value): return None
    return value.isoformat() if hasattr(value,'isoformat') else value

def _table(name, rows, styles=None, merges=None, widths=None):
    nonempty=[(n,row) for n,row in enumerate(rows) if any(v is not None and v!='' for v in row)]
    candidates=nonempty[:20]
    # A title row has few cells; column labels are predominantly strings.
    headers=[(n,row) for n,row in candidates if sum(isinstance(v,str) and bool(v.strip()) for v in row)>=max(1,.8*sum(v is not None and v!='' for v in row))]
    header,row=max(headers or candidates,key=lambda item:sum(v is not None and v!='' for v in item[1]),default=(0,[]))
    columns=[str(v) if v is not None and v!='' else f'Column {i+1}' for i,v in enumerate(row)]
    columns=[v if columns.count(v)==1 else f'{v} [{i+1}]' for i,v in enumerate(columns)]
    return {'name':name,'columns':columns,'header_row':header+1,'rows':rows,
            'records':[{columns[i]:scalar(v) for i,v in enumerate(row[:len(columns)])} for n,row in nonempty if n>header],
            'record_rows':[n+1 for n,row in nonempty if n>header],
            'styles':styles or {},'merges':merges or [],'widths':widths or {}}

@lru_cache(maxsize=128)
def _load(path,mtime,size):
    path=Path(path)
    if size>20*1024*1024: raise WorkbenchError('file_too_large','Preview exceeds file limit')
    if path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        with ZipFile(path) as archive:
            if len(archive.infolist())>4000 or sum(i.file_size for i in archive.infolist())>80*1024*1024:
                raise WorkbenchError('file_too_large','Spreadsheet expansion exceeds preview limit')
        book=load_workbook(path,data_only=True,read_only=False,keep_links=False)
        try:
            if len(book.worksheets)>200 or sum(len(s._cells) for s in book)>100000:
                raise WorkbenchError('file_too_large','Spreadsheet exceeds cell or sheet limit')
            tables=[]
            for sheet in book:
                # Bound stored coordinates; do not allocate attacker-controlled dimensions.
                cells=list(sheet._cells.values())
                maxrow=max((c.row for c in cells),default=0); maxcol=max((c.column for c in cells),default=0)
                if maxrow>20000 or maxcol>256 or maxrow*maxcol>200000:
                    raise WorkbenchError('file_too_large','Worksheet preview exceeds 20,000 rows or 256 columns')
                rows=[[None]*maxcol for _ in range(maxrow)]; styles={}
                for c in cells:
                    rows[c.row-1][c.column-1]=scalar(c.value)
                    styles[c.coordinate]={'bold':bool(c.font.bold),'italic':bool(c.font.italic),
                        'color':c.font.color.rgb[-6:] if c.font.color and c.font.color.type=='rgb' else None,
                        'fill':c.fill.fgColor.rgb[-6:] if c.fill.patternType=='solid' and c.fill.fgColor.type=='rgb' else None,
                        'align':c.alignment.horizontal or 'general','format':c.number_format}
                tables.append(_table(sheet.title,rows,styles,[str(m) for m in sheet.merged_cells.ranges],
                    {k:v.width for k,v in sheet.column_dimensions.items()}))
            return tables
        finally: book.close()
    data=path.read_bytes()
    try: text=data.decode('utf-16' if data.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig')
    except UnicodeError as exc: raise WorkbenchError('parse_failed','Table must use UTF-8 or UTF-16 text') from exc
    if path.suffix.lower() in {'.csv','.tsv'}:
        rows=list(csv.reader(text.splitlines(),delimiter='\t' if path.suffix.lower()=='.tsv' else ','))
        if len(rows)>20000 or any(len(r)>256 for r in rows): raise WorkbenchError('file_too_large','Table exceeds preview limits')
        return [_table(path.stem,rows)]
    values=[json.loads(line) for line in text.splitlines() if line.strip()] if path.suffix.lower()=='.jsonl' else json.loads(text)
    if isinstance(values,dict):
        groups={k:v for k,v in values.items() if isinstance(v,list) and v and all(isinstance(x,dict) for x in v)}
        if not groups: groups={path.stem:[values]}
    else: groups={path.stem:values}
    tables=[]
    for name,records in groups.items():
        if not isinstance(records,list) or len(records)>20000: raise WorkbenchError('file_too_large','Table exceeds record limit')
        columns=list(dict.fromkeys(k for r in records if isinstance(r,dict) for k in r))
        if len(columns)>256: raise WorkbenchError('file_too_large','Table exceeds column limit')
        rows=[columns]+[[r.get(k) if isinstance(r.get(k),(str,int,float,bool,type(None))) else json.dumps(r.get(k),ensure_ascii=False) for k in columns] for r in records if isinstance(r,dict)]
        tables.append(_table(str(name),rows))
    return tables

def load_tables(path):
    path=Path(path); stat=path.stat()
    return _load(str(path.resolve()),stat.st_mtime_ns,stat.st_size)

def preview(path,sheet=0,offset=0,limit=100):
    from openpyxl.utils import get_column_letter
    tables=load_tables(path)
    if not 0<=sheet<len(tables) or offset<0 or not 1<=limit<=200: raise WorkbenchError('invalid_request','Invalid worksheet or row range')
    table=tables[sheet]; rows=table['rows']
    return {'sheets':[{'name':t['name'],'row_count':len(t['rows']),'column_count':max(map(len,t['rows']),default=0)} for t in tables],
        'sheet':sheet,'name':table['name'],'offset':offset,'total':len(rows),'columns':table['columns'],
        'merges':table['merges'],'widths':table['widths'],
        'rows':[{'number':i+1,'cells':[{'coordinate':f'{get_column_letter(j+1)}{i+1}','value':value,
            **table['styles'].get(f'{get_column_letter(j+1)}{i+1}',{})} for j,value in enumerate(row)]} for i,row in enumerate(rows[offset:offset+limit],offset)]}

def find_cells(path,q,sheet=None,match_case=False,whole_cell=False,offset=0,limit=200):
    from openpyxl.utils import get_column_letter
    if not isinstance(q,str) or not q.strip() or len(q)>256 or offset<0 or not 1<=limit<=200:
        raise WorkbenchError('invalid_request','Enter a search value under 256 characters and a valid result range')
    tables=load_tables(path)
    if sheet is not None and not 0<=sheet<len(tables): raise WorkbenchError('invalid_request','Unknown worksheet')
    needle=q if match_case else q.casefold()
    matches=[];total=0
    for index,table in enumerate(tables):
        if sheet is not None and index!=sheet: continue
        for row,values in enumerate(table['rows'],1):
            for column,value in enumerate(values,1):
                if value is None: continue
                text=str(value);candidate=text if match_case else text.casefold()
                if (candidate==needle if whole_cell else needle in candidate):
                    if offset<=total<offset+limit:
                        matches.append({'sheet':index,'name':table['name'],'row':row,'column':column,
                            'coordinate':f'{get_column_letter(column)}{row}','value':value})
                    total+=1
    return {'matches':matches,'total':total,'offset':offset,'limit':limit}

def literal_entity_coverage(table,question):
    """Bounded observations of literal request terms across complete records.

    Matches are source facts, not inferred entities or routing decisions. This
    precedes semantic applicability so a two-row preview cannot imply absence.
    """
    terms=list(dict.fromkeys(token for token in re.findall(r'\b[\w-]+\b',question)
        if 2<=len(token)<=80 and any(char.isalpha() for char in token)))
    selected=terms[:40]
    matches=[]
    for column in table['columns']:
        for term in selected:
            count=0;examples=[]
            for row in table['records']:
                value=str(row.get(column) or '').strip()
                if term.casefold() in value.casefold():
                    count+=1
                    if len(examples)<2 and value[:160] not in examples:examples.append(value[:160])
            if count:matches.append({'literal':term,'column':column,'matching_records':count,'examples':examples})
    bounded=[];size=0
    for match in matches:
        length=len(json.dumps(match,ensure_ascii=True))+2
        if size+length>1900:break
        bounded.append(match);size+=length
    return {'scanned_rows':len(table['records']),'matches':bounded,
        'truncated':len(terms)>40 or len(bounded)<len(matches)}


def population_binding_catalog(available,entities):
    """Observed source candidates for a separate semantic identity-field pass."""
    candidates=[]
    if not entities:return candidates
    for identifier,(_,table) in available.items():
        for column in table['columns']:
            matches={str(entity):[str(row.get(column)).strip() for row in table['records']
                if str(entity).casefold() in str(row.get(column)).strip().casefold()] for entity in entities}
            if all(matches.values()):
                candidates.append({'table':identifier,'column':column,
                    'matching_records':{entity:len(values) for entity,values in matches.items()},
                    'examples':{entity:[value[:160] for value in list(dict.fromkeys(values))[:2]] for entity,values in matches.items()}})
    return candidates


def compile_aggregate_request(available, question):
    """Bind explicit compound aggregates to actual fields before generation.

    This handles mathematical wording only; ambiguous field roles continue
    through the existing semantic contract. No source values are invented.
    """
    candidates=[]
    for identifier,(_,table) in available.items():
        metrics=[];groups=[]
        if re.search(r'\b(?:number|count)\s+of\s+(?:records|rows)|\brecord\s+counts?\b',question,re.I):
            metrics.append({'operation':'count','column':None})
        for column in table['columns']:
            name=re.escape(column)
            if re.search(r'\b(?:per|by)\s+'+name+r'\b|(?<!\w)'+name+r'\s+counts?\b',question,re.I):groups.append(column)
            if re.search(r'\bcohort\s+counts?\b',question,re.I) and re.fullmatch(r'group|cohort|team|category|department',column,re.I):groups.append(column)
            for op,words in [('sum','sum|total|totals'),('average','average|mean'),('min','minimum|min'),('max','maximum|max')]:
                if re.search(r'\b(?:'+words+r')\s+(?:of\s+)?(?:numeric\s+)?'+name+r'\b|(?<!\w)'+name+r'\s+(?:'+words+r')\b',question,re.I):
                    metrics.append({'operation':op,'column':column})
        groups=list(dict.fromkeys(groups))
        if groups and not any(m['operation']=='count' for m in metrics):metrics.insert(0,{'operation':'count','column':None})
        if groups or len(metrics)>1:
            filters=[]
            for column in groups:
                named=[value for value in {str(row.get(column)) for row in table['records']} if re.search(r'(?<!\w)'+re.escape(value)+r'(?!\w)',question,re.I)]
                if named:filters.append({'column':column,'operator':'in','value':named})
            candidates.append({'table':identifier,'operation':'aggregate','columns':[],'filters':filters,'metrics':metrics,'group_by':groups})
    return candidates[0] if len(candidates)==1 else None


def compile_outcome_query(intent,available,question,history=None):
    """Bind a semantic cohort/outcome measure to complete source records.

    Exact identities and prefixes are resolved from actual records, and every
    compatible assessment is included without model-authored arithmetic.
    """
    literal_criteria={}
    for identifier,(_,table) in available.items():
        predicates=[]
        for field in table['columns']:
            # Only source-bound, explicitly named field comparisons qualify.
            pattern=r'(?<!\w)'+re.escape(field)+r'\s*(?:is\s+)?(at least|at most|above|below|>=|<=|>|<|equals?|=)\s*(-?\d+(?:\.\d+)?)'
            match=re.search(pattern,question,re.I)
            if match:
                operator={'at least':'gte','at most':'lte','above':'gt','below':'lt','>=':'gte','<=':'lte','>':'gt','<':'lt','equal':'eq','equals':'eq','=':'eq'}[match[1].lower()]
                predicates.append({'column':field,'operator':operator,'value':float(match[2])})
            label=re.search(r'(?<!\w)'+re.escape(field)+r'\s*(?:is|=|equals?|of)?\s*(PASS|FAIL|PASSED|FAILED)(?!\w)',question,re.I)
            if label:
                values={str(row.get(field)).strip() for row in table['records'] if str(row.get(field)).strip().casefold()==label[1].casefold()}
                if values:predicates.append({'column':field,'operator':'in','value':sorted(values)})
        if predicates:literal_criteria[identifier]=predicates
    outcomes=list(dict.fromkeys(intent.get('result_values',[])))
    polarity=intent.get('outcome')
    canonical={'pass':{'pass','passed'},'fail':{'fail','failed'}}
    if polarity in canonical:
        # The model chooses meaning, never an arbitrary union of outcome values.
        outcomes=sorted({str(row.get(column)).strip() for _,table in available.values()
            for row in table['records'] for column in table['columns']
            if str(row.get(column)).strip().casefold() in canonical[polarity]})
    entities=list(dict.fromkeys(intent.get('entity_values',[])))
    resolved=intent.get('_identity_resolution')
    user_history='\n'.join(item for item in (history or []) if isinstance(item,str) and item.lower().startswith('user:'))
    # A selected identity present only in prior user requests is inherently
    # inherited context; an inconsistent followup flag must not discard it.
    inherited=bool(entities) and all(str(value).casefold() in user_history.casefold() for value in entities) and not all(str(value).casefold() in question.casefold() for value in entities)
    if inherited:intent={**intent,'followup':True}
    explicit_scores=[field for field in intent.get('score_columns',[]) if len(re.sub(r'[^a-z0-9]','',field.casefold()))>=3 and
        re.search(r'(?<!\w)'+r'[\W_]*'.join(re.findall(r'[a-z]+|\d+',field.casefold()))+r'(?!\w)',question.casefold())]
    if explicit_scores:
        intent={**intent,'score_columns':explicit_scores,'scope':'one','assessments':[
            table['name'] for _,table in available.values() if all(field in table['columns'] for field in explicit_scores)]}
    grouped=bool(intent.get('group_entities'))
    if grouped and len(entities)<2:
        raise WorkbenchError('invalid_query','A cohort comparison requires every requested group, not a single inherited cohort')
    context=question+'\n'+'\n'.join(history or []) if intent.get('followup') else question
    # Bind literal identifiers/prefixes even when the semantic model omits
    # them. Only tokens present in one real source column qualify; ordinary
    # words and standalone thresholds cannot silently become cohort filters.
    user_context=question
    if intent.get('followup'):
        user_context+='\n'+'\n'.join(item for item in (history or []) if not item.lower().startswith('assistant:'))
    inferred=[]
    # Only the current request can introduce a literal source-bound cohort.
    # Historical identifiers can belong to unrelated completed requests; the
    # semantic resolver selects which antecedent entities remain active.
    for token in dict.fromkeys(re.findall(r'\b[A-Za-z0-9]+\b',question)):
        if not (any(c.isalpha() for c in token) and any(c.isdigit() for c in token)): continue
        candidates={column for _,table in available.values() for column in table['columns']
            if any(str(row.get(column)).strip().casefold().startswith(token.casefold()) for row in table['records'])}
        if len(candidates)==1: inferred.append(token)
    if resolved:
        inferred=resolved['canonical']
        context+='\n'+' '.join(inferred)
    if inferred:
        if entities and {v.casefold() for v in entities}!={v.casefold() for v in inferred}:
            raise WorkbenchError('invalid_query','Planned cohort does not match the source-bound user identifiers. Expected '+repr(inferred)+'; proposed '+repr(entities))
        entities=inferred
    if any(str(value).casefold() not in context.casefold() for value in entities):
        raise WorkbenchError('invalid_query','Cohort identities must be literal names from the current question or explicitly referenced history; do not expand source prefixes. Proposed '+repr(entities))
    threshold=intent.get('threshold')
    applied_operator=intent.get('threshold_operator')
    if threshold is not None:
        if not isinstance(threshold,(int,float)) or not math.isfinite(threshold) or intent.get('threshold_operator') not in {'gte','gt','lt','lte'}:
            raise WorkbenchError('invalid_query','Invalid numeric outcome rule')
        supplied=[float(v) for v in re.findall(r'(?<![\w.])-?\d+(?:\.\d+)?(?!\w|\.\d)',user_context)]
        if threshold not in supplied:
            # A proposed numeric rule is not evidence. Explicit source outcomes
            # can answer the measure without that rule; retain the population
            # and assessment scope and bind those recorded outcomes instead.
            if outcomes:
                threshold=None
                applied_operator=None
            else:
                raise WorkbenchError('needs_input','The available records have no explicit outcome labels or user-supplied numeric rule. Supply the rule to calculate this outcome.')
        if threshold is not None and not intent.get('score_columns'): raise WorkbenchError('needs_input','Specify the assessment score columns for the numeric rule')
        # Keep the user's literal rule separate from the requested numerator.
        # Derive its complement once, at the execution boundary, rather than
        # mutating conversational intent that a later follow-up may reuse.
        rule_outcome=intent.get('rule_outcome')
        if threshold is not None and rule_outcome in canonical:
            if polarity not in canonical:
                raise WorkbenchError('needs_input','Specify whether the numeric rule measures passing or failing records')
            if rule_outcome!=polarity:
                applied_operator={'lt':'gte','lte':'gt','gt':'lte','gte':'lt'}[applied_operator]
    queries=[]; documents=set(); unavailable=[]
    for identifier,(doc,table) in available.items():
        if intent.get('scope')!='all' and intent.get('assessments') and table['name'] not in intent['assessments']: continue
        criteria=[]
        for column in table['columns']:
            values={str(row.get(column)).strip().casefold() for row in table['records']}
            if threshold is not None and column in intent['score_columns']:
                if any(numeric_value(row.get(column)) is not None for row in table['records']):
                    criteria.append({'column':column,'operator':applied_operator,'value':threshold})
            elif threshold is None and outcomes:
                bound=[value for value in outcomes if str(value).casefold() in values]
                if bound and (polarity in canonical or len(bound)==len(outcomes)):
                    criteria.append({'column':column,'operator':'in','value':bound})
        if identifier in literal_criteria:
            criteria=literal_criteria[identifier]
        if not criteria:
            numeric_fields=[field for field in table['columns'] if field!=intent.get('entity_column') and
                any(numeric_value(row.get(field)) is not None for row in table['records'])]
            if threshold is not None and intent.get('scope')=='all' and len(numeric_fields)>1:
                raise WorkbenchError('invalid_query','The supplied numeric rule has incomplete assessment coverage. '
                    'Bind relevant score_columns from the consolidated source '+table['name']+': '+repr(numeric_fields)+
                    '. Do not request the threshold again or omit these measurements from an all-assessments report.')
            unavailable.append(table['name']);continue
        filters=[]
        if entities:
            discovered=exact_identity_columns(table,entities)
            complete=[field for field,counts in discovered.items() if all(counts.values())]
            proposed=intent.get('entity_column')
            if len(complete)==1:
                intent={**intent,'entity_column':complete[0]}
            elif len(complete)>1 and proposed not in complete:
                raise WorkbenchError('needs_input','Exact identities occur in multiple fields in '+table['name']+'. Choose from '+repr(complete))
            candidates=[]
            bindings={}
            for column in table['columns']:
                values={identity_value(row.get(column)) for row in table['records']}
                exact=all(identity_value(entity) in values for entity in entities)
                partial=all(any(identity_value(entity) in value for value in values) for entity in entities)
                if exact or partial:
                    bindings[column]={str(entity):sum(str(entity).casefold() in str(row.get(column)).strip().casefold() for row in table['records']) for entity in entities}
                if intent.get('entity_column') is not None and column!=intent['entity_column']:continue
                # A complete identifier cannot be established by a substring.
                explicit_id=any(any(c.isalpha() for c in str(e)) and any(c.isdigit() for c in str(e)) for e in entities)
                if exact or partial and not explicit_id: candidates.append((column,exact))
            exact=[candidate for candidate in candidates if candidate[1]]
            candidates=exact or candidates
            if len(candidates)!=1:
                raise WorkbenchError('needs_input','The cohort does not resolve to one identity column in '+table['name']+
                    '. Proposed population values: '+repr(entities)+'. An assessment name is not a population; measurement follow-ups retain the prior requested population. Choose entity_column from the source-bound fields and observed matching record counts: '+repr(bindings)+
                    '. Proposed field: '+repr(intent.get('entity_column')))
            column,is_exact=candidates[0]
            if is_exact: filters=[{'column':column,'operator':'in','value':entities}]
            elif len(entities)==1 or grouped: filters=[{'column':column,'operator':'contains','value':entities[0]}]
            else: raise WorkbenchError('needs_input','Use complete identities or one cohort prefix for percentages')
        populations=entities if grouped else [None]
        for population in populations:
            population_filters=([{'column':column,'operator':'eq' if is_exact else 'contains','value':population}]
                                if population is not None else filters)
            queries.append({'table':identifier,'columns':[],'filters':population_filters,'criteria':criteria,'requested_outcome':polarity,
                            **({'criteria_mode':'all'} if identifier in literal_criteria else {}),
                            **({'cohort':population} if population is not None else {})})
        documents.add(doc['document_id'])
    if len(documents)>1: raise WorkbenchError('needs_input','Connect one results document for this multi-assessment percentage')
    if not queries: raise WorkbenchError('needs_input','No explicit result values were found. Specify the numeric pass rule before calculating percentages.')
    operation=intent.get('operation','percentage')
    # A consolidated score table already covers each assessment. Avoid counting
    # its duplicate weekly views as additional tests.
    consolidated=[query for query in queries if len(query['criteria'])>1]
    if threshold is not None and len({query['table'] for query in consolidated})==1:
        queries=consolidated;unavailable=[]
    first,*rest=queries
    return {**first,'operation':operation,'additional_queries':rest,'unavailable_assessments':unavailable}


def bind_selection_request(plan,table,question):
    """Retain literal identities, including missing siblings of source IDs."""
    question=unicodedata.normalize('NFC',question)
    filters=plan.get('filters',[])
    rule=re.search(r'\b(\d+(?:\.\d+)?)\s+is\s+(?:the\s+)?pass(?:ing)?\s+(?:mark|score|threshold)\b|\bpass(?:ing)?\s+(?:mark|score|threshold)\s*(?:is|=|:)\s*(\d+(?:\.\d+)?)',question,re.I)
    tokens=list(dict.fromkeys(re.findall(r'(?<!\w)[\w]+(?:[.-][\w]+)*(?!\w)',question)))
    literal_ids=[token for token in tokens if any(c.isalpha() for c in token) and any(c.isdigit() for c in token)]
    literal_ids.extend(str(value) for item in filters if item.get('column') not in table['columns'] and item.get('operator') in {'eq','in'}
                       for value in (item['value'] if isinstance(item.get('value'),list) else [item.get('value')])
                       if isinstance(value,str) and re.search(r'(?<!\w)'+re.escape(value)+r'(?!\w)',question) and value not in literal_ids)
    discovered=exact_identity_columns(table,literal_ids) if literal_ids else {}
    if discovered:
        best=max(sum(count>0 for count in counts.values()) for counts in discovered.values())
        fields=[field for field,counts in discovered.items() if sum(count>0 for count in counts.values())==best]
        if len(fields)!=1:
            raise WorkbenchError('needs_input','Exact identifiers occur in multiple source fields. Choose from '+repr(fields))
        field=fields[0]
        found=[token for token in literal_ids if discovered[field][token]]
        families={re.sub(r'\d+$','',identity_value(token)) for token in found}
        entities=[token for token in literal_ids if token in found or re.sub(r'\d+$','',identity_value(token)) in families]
        proposed_fields={f['column'] for f in filters if f.get('operator') in {'eq','in','contains'} and
            any(identity_value(v) in {identity_value(e) for e in entities} for v in (f['value'] if isinstance(f.get('value'),list) else [f.get('value')]))}
        filters=[f for f in filters if f['column'] not in proposed_fields and f['column']!=field]
        unresolved=re.findall(r'\band\s+(\d+)\b',question,re.I)
        return {**plan,'filters':filters+[{'column':field,'operator':'in','value':entities}],
                'requested_entities':entities,'identity_column':field,'entity_bindings':{},
                **({'unresolved_identifiers':unresolved} if unresolved else {}),
                **({'pass_threshold':float(rule[1] or rule[2])} if rule else {}),
                'comparison_requested':bool(re.search(r'\b(?:compare|comparison|better|best|highest|lowest|rank)\b',question,re.I))}
    # Preview examples are source data, never authority to choose the user's population.
    for item in filters:
        if item.get('operator') not in {'eq','in','contains'}:continue
        values=item.get('value',[])
        for value in values if isinstance(values,list) else [values]:
            if not isinstance(value,str) or not any(c.isalpha() for c in value) or not any(c.isdigit() for c in value):continue
            if exact_identity_columns(table,[value]) and identity_value(value) not in {identity_value(v) for v in literal_ids}:
                raise WorkbenchError('needs_input','The proposed table query introduced source identifiers absent from your request. Specify the entity or relevant source; no table values were presented.')
    if plan.get('operation')!='select':return plan
    identity=next((f['column'] for f in filters if f.get('operator') in {'eq','in'}),None)
    # A business filter can precede the identity filter. Bind literal source
    # IDs to their unique filtered column rather than depending on list order.
    id_tokens=[token.casefold() for token in re.findall(r'\b[\w-]+\b',question)
        if any(c.isalpha() for c in token) and re.search(r'\d$',token)]
    candidates=[]
    for column in table['columns']:
        values={str(r.get(column,'')).strip().casefold() for r in table['records']}
        families={re.sub(r'\d+$','',v) for v in values if re.fullmatch(r'[\w-]*[A-Za-z][\w-]*\d+',v)}
        if any(token in values or re.sub(r'\d+$','',token) in families for token in id_tokens):
            candidates.append(column)
    if len(candidates)==1:identity=candidates[0]
    if identity not in table['columns']:return plan
    values={str(r.get(identity,'')).strip().casefold() for r in table['records']}
    families={re.sub(r'\d+$','',v) for v in values if re.fullmatch(r'[\w-]*[A-Za-z][\w-]*\d+',v)}
    tokens=list(dict.fromkeys(re.findall(r'\b[\w-]+\b',question)))
    entities=[token for token in tokens if token.casefold() in values or (re.sub(r'\d+$','',token.casefold()) in families and re.search(r'\d$',token))]
    for item in filters:
        if item.get('column')!=identity:continue
        raw=item.get('value',[])
        for value in raw if isinstance(raw,list) else [raw]:
            if isinstance(value,str) and value.casefold() in question.casefold() and value.casefold() not in {v.casefold() for v in entities}:entities.append(value)
    bindings={}
    if entities:
        filters=[f for f in filters if f.get('column')!=identity]+[{'column':identity,'operator':'in','value':entities}]
    rule=re.search(r'\b(\d+(?:\.\d+)?)\s+is\s+(?:the\s+)?pass(?:ing)?\s+(?:mark|score|threshold)\b|\bpass(?:ing)?\s+(?:mark|score|threshold)\s*(?:is|=|:)\s*(\d+(?:\.\d+)?)',question,re.I)
    return {**plan,'filters':filters,'requested_entities':entities,'identity_column':identity,'entity_bindings':bindings,
            **({'pass_threshold':float(rule[1] or rule[2])} if rule else {}),
            'comparison_requested':bool(re.search(r'\b(?:compare|comparison|better|best|highest|lowest|rank)\b',question,re.I))}

def execute_query(table,plan):
    operation=plan.get('operation'); columns=plan.get('columns',[]); filters=plan.get('filters',[])
    measurement_operation=operation in {'overall_average','measurement_rank','measurement_missing','measurement_count','measurement_sum'}
    if operation not in {'select','count','sum','average','min','max','percentage','aggregate','overall_average','measurement_rank','measurement_missing','measurement_count','measurement_sum'} or not isinstance(columns,list) or not isinstance(filters,list) or len(columns)>(256 if measurement_operation or operation=='select' else 12) or len(filters)>8:
        raise WorkbenchError('invalid_query','Unsupported table query')
    if any(c not in table['columns'] for c in columns): raise WorkbenchError('invalid_query','Query refers to an unknown column')
    for f in filters:
        if not isinstance(f,dict) or f.get('column') not in table['columns'] or f.get('operator') not in {'eq','ne','contains','in','gt','gte','lt','lte'}:
            raise WorkbenchError('invalid_query','Query has an unknown column or operator')
        value=f.get('value')
        if (f['operator']=='in' and (not isinstance(value,list) or not 1<=len(value)<=40 or any(not isinstance(v,(str,int,float)) for v in value)) or
            f['operator']!='in' and not isinstance(value,(str,int,float))):
            raise WorkbenchError('invalid_query','Query filter has an invalid value or alternative list')
    equalities={}
    for f in filters:
        if f['operator']=='eq':
            value=str(f['value']).strip().casefold()
            if f['column'] in equalities and equalities[f['column']]!=value:
                raise WorkbenchError('invalid_query','Query has mutually exclusive equalities; use an alternative-value filter')
            equalities[f['column']]=value
    def match(row,f):
        a=row.get(f['column']); b=f['value']; op=f['operator']
        if op=='in': return identity_value(a) in {identity_value(v) for v in b}
        if op in {'eq','ne','contains'}:
            if op in {'eq','ne'} and isinstance(b,(int,float)) and not isinstance(b,bool):
                left=numeric_value(a);right=numeric_value(b)
                if left is None or right is None:return False
                return left==right if op=='eq' else left!=right
            a=identity_value(a); b=identity_value(b)
            return a==b if op=='eq' else a!=b if op=='ne' else b in a
        a=numeric_value(a);b=numeric_value(b)
        if a is None or b is None:return False
        return {'gt':a>b,'gte':a>=b,'lt':a<b,'lte':a<=b}[op]
    rows=[r for r in table['records'] if all(match(r,f) for f in filters)]
    ordering=plan.get('order_by',[])
    limit=plan.get('limit')
    if not isinstance(ordering,list) or len(ordering)>3 or any(not isinstance(o,dict) or o.get('column') not in table['columns'] or o.get('direction') not in {'asc','desc'} for o in ordering):
        raise WorkbenchError('invalid_query','Invalid ordering column or direction')
    if limit is not None and (isinstance(limit,bool) or not isinstance(limit,int) or not 1<=limit<=40):
        raise WorkbenchError('invalid_query','Limit must be an integer from 1 to 40')
    if ordering and operation not in {'select','aggregate'}:
        raise WorkbenchError('invalid_query','Ordering requires selection or grouped aggregation')
    for order in reversed(ordering if operation=='select' else []):
        column=order['column']
        present=[r for r in rows if r.get(column) is not None]
        missing=[r for r in rows if r.get(column) is None]
        numeric=all(numeric_value(r.get(column)) is not None for r in present)
        rows=sorted(present,key=lambda r:numeric_value(r[column]) if numeric else str(r[column]).casefold(),reverse=order['direction']=='desc')+missing
    result={'table':table['name'],'operation':operation,'scanned_rows':len(table['records']),'matched_rows':len(rows),'columns':columns,'filters':filters}
    row_numbers=table.get('record_rows',list(range(table.get('header_row',1)+1,table.get('header_row',1)+1+len(table['records']))))
    selected_rows={id(row) for row in rows}
    result['record_locators']=[{'sheet':table['name'],'row':number} for number,row in zip(row_numbers,table['records']) if id(row) in selected_rows]
    if plan.get('unresolved_identifiers'):result['unresolved_identifiers']=plan['unresolved_identifiers']
    if plan.get('requested_outcome') in {'pass','fail'}:
        result['requested_outcome']=plan['requested_outcome']
    if measurement_operation:
        from rag.analytics import execute_measurements
        return execute_measurements(table,plan,rows,result)
    if operation=='aggregate':
        metrics=plan.get('metrics',[]);groups=plan.get('group_by',[])
        if not isinstance(metrics,list) or not 1<=len(metrics)<=12 or not isinstance(groups,list) or len(groups)>3 or any(c not in table['columns'] for c in groups):
            raise WorkbenchError('invalid_query','Invalid aggregate contract')
        def calculate(records):
            values=[]
            for metric in metrics:
                op=metric.get('operation');column=metric.get('column')
                if op not in {'count','sum','average','min','max'} or (column not in table['columns'] and not (op=='count' and column is None)):
                    raise WorkbenchError('invalid_query','Unknown aggregate metric/field')
                measured=execute_query({**table,'records':records},{'operation':op,'columns':[column] if column else [],'filters':[]})
                values.append({'operation':op,'column':column, 'value':measured.get('value'),'numeric_rows':measured.get('numeric_rows')})
            return values
        result['metrics']=calculate(rows);result['group_by']=groups;result['groups']=[]
        grouped={}
        for row in rows:
            key=tuple(row.get(c) for c in groups)
            grouped.setdefault(key,[]).append(row)
        if groups:
            if len(grouped)>200:raise WorkbenchError('resource_limit','Too many aggregate groups')
            result['groups']=[{'keys':dict(zip(groups,key)),'metrics':calculate(records)} for key,records in grouped.items()]
            for order in reversed(ordering):
                column=order['column']
                if column in groups:
                    key=lambda group:str(group['keys'][column]).casefold()
                else:
                    indexes=[i for i,m in enumerate(metrics) if m['column']==column]
                    if len(indexes)!=1:
                        raise WorkbenchError('invalid_query','Grouped ordering must identify one aggregate metric')
                    index=indexes[0]
                    key=lambda group:group['metrics'][index]['value']
                present=[g for g in result['groups'] if key(g) is not None]
                missing=[g for g in result['groups'] if key(g) is None]
                result['groups']=sorted(present,key=key,reverse=order['direction']=='desc')+missing
            if limit is not None:result['groups']=result['groups'][:limit]
            result.update(order_by=ordering,limit=limit)
        elif ordering or limit is not None:
            raise WorkbenchError('invalid_query','Aggregate ordering and limit require grouping')
    elif operation=='select':
        bound=limit or 40
        result.update(records=[{c:r.get(c) for c in columns or table['columns']} for r in rows[:bound]],truncated=limit is None and len(rows)>40,order_by=ordering,limit=limit)
        from openpyxl.utils import get_column_letter
        locators={id(row):number for number,row in zip(row_numbers,table['records'])}
        result['cell_locators']=[{'sheet':table['name'],'row':locators[id(row)],
            'cells':{field:get_column_letter(table['columns'].index(field)+1)+str(locators[id(row)]) for field in columns or table['columns']}}
            for row in rows[:bound]]
        identity=plan.get('identity_column')
        if identity:
            present={identity_value(r.get(identity,'')) for r in rows}
            result['missing_entities']=[v for v in plan.get('requested_entities',[]) if identity_value(v) not in present]
            result['entity_bindings']=plan.get('entity_bindings',{})
        if plan.get('comparison_requested') and not result['truncated'] and identity:
            fields=[c for c in columns or table['columns'] if c!=identity and any(numeric_value(r.get(c)) is not None for r in rows)]
            winners=[];minima=[]
            for field in fields:
                numeric=[r for r in rows if numeric_value(r.get(field)) is not None]
                if numeric:
                    maximum=max(numeric_value(r[field]) for r in numeric)
                    winners.append({'column':field,'maximum':int(maximum) if maximum.is_integer() else maximum,'identities':[str(r.get(identity)) for r in numeric if numeric_value(r[field])==maximum]})
                    minimum=min(numeric_value(r[field]) for r in numeric)
                    minima.append({'column':field,'minimum':int(minimum) if minimum.is_integer() else minimum,'identities':[str(r.get(identity)) for r in numeric if numeric_value(r[field])==minimum]})
            result['comparison']={'identity_column':identity,'field_maxima':winners,'field_minima':minima}
            if 'pass_threshold' in plan:
                threshold=plan['pass_threshold']
                result['comparison']['pass_threshold']=threshold
                result['comparison']['pass_counts']=[{'identity':str(r.get(identity)),
                    'passed':sum(numeric_value(r.get(c)) is not None and numeric_value(r[c])>=threshold for c in fields),
                    'assessments':len(fields)} for r in rows]
    elif operation in {'count','percentage'} and plan.get('criteria_mode')=='all':
        criteria=plan.get('criteria',[])
        numerator=execute_query({**table,'records':rows},{'operation':'count','columns':[],'filters':criteria})['value']
        result.update(value=numerator,total=len(rows),criteria=criteria,criteria_mode='all')
        result['summaries']=[{'column':' AND '.join(c['column'] for c in criteria),'value':numerator,'passed':numerator,'total':len(rows),
                              **({'percentage':round(100*numerator/len(rows),2) if rows else 0} if operation=='percentage' else {})}]
    elif operation=='count':
        criteria=plan.get('criteria',[])
        if criteria:
            if not isinstance(criteria,list) or len(criteria)>12: raise WorkbenchError('invalid_query','Too many result count criteria')
            result['summaries']=[{'column':criterion.get('column'),'value':execute_query({**table,'records':rows},
                {'operation':'count','columns':[],'filters':[criterion]})['value'],'total':len(rows),'criterion':criterion} for criterion in criteria]
            if len(criteria)==1: result['value']=result['summaries'][0]['value']
        else: result['value']=len(rows)
    elif operation=='percentage':
        criteria=plan.get('criteria',[])
        if not criteria or len(criteria)>12:
            raise WorkbenchError('invalid_query','A percentage requires an explicit result value or a source/user supplied threshold for each assessment')
        # Reuse the same validated comparison contract for numerator predicates.
        summaries=[]
        for criterion in criteria:
            if criterion.get('operator')=='in':
                existing={str(row.get(criterion.get('column'))).strip().casefold() for row in table['records']}
                values=criterion.get('value',[])
                if isinstance(values,list) and any(str(value).strip().casefold() not in existing for value in values):
                    raise WorkbenchError('invalid_query','Percentage criterion values do not exist in '+table['name']+'. Choose a result column from categorical_values in the catalog; numeric mark columns do not contain PASS.')
            probe=execute_query({**table,'records':rows},{'operation':'count','columns':[], 'filters':[criterion]})
            summaries.append({'column':criterion['column'],'passed':probe['value'],'total':len(rows),
                'percentage':round(100*probe['value']/len(rows),2) if rows else None,'criterion':criterion})
        result['summaries']=summaries
    else:
        if not columns: raise WorkbenchError('invalid_query','Numeric aggregation requires at least one column')
        if len(columns)>1:
            result['summaries']=[{'column':column, **{k:v for k,v in execute_query(table,{**plan,'columns':[column]}).items() if k in {'value','numeric_rows'}}} for column in columns]
            return result
        numbers=[]
        for row in rows:
            value=row.get(columns[0])
            if isinstance(value,bool): continue
            try: number=float(value)
            except (TypeError,ValueError): continue
            if math.isfinite(number): numbers.append(number)
        result['numeric_rows']=len(numbers)
        result['value']=sum(numbers) if operation=='sum' else sum(numbers)/len(numbers) if operation=='average' and numbers else min(numbers) if operation=='min' and numbers else max(numbers) if operation=='max' and numbers else None
    return result
