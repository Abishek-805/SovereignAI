import json,time,hashlib,urllib.request
from scripts.full_qa_inventory import http,save,OUT
from io import BytesIO
from docx import Document
from openpyxl import load_workbook
cases=json.loads((OUT/'document-audit.json').read_text(encoding='utf-8'))['cases'];ids={c['name']:c.get('document_id') for c in cases}
results=[]
for case,goal,did in [('AT-15','Search the selected document and create a Word report in Knowledge named qa_source_report.docx. Include only the number of repair bays, the supervisor name, and the external network rule, with source attribution.',ids['qa-source.txt']),('AT-16','Analyze the selected spreadsheet and create an Excel summary in Knowledge named qa_source_summary.xlsx. One sheet Summary, columns Group and Count. Count all records per Group including blank Score records. Expected groups must come from the source.',ids['qa-source.xlsx'])]:
 admission=http('/agent/jobs','POST',{'goal':goal,'document_ids':[did],'history':[]});r={'id':case,'goal':goal,'source_id':did,'admission':admission};save('agent/'+case+'.json',r)
 if admission['status']==200:
  jid=admission['body']['job_id'];start=time.monotonic()
  for _ in range(600):
   final=http('/coding/jobs/'+jid)['body']
   if final.get('state')!='running':break
   time.sleep(1)
  r['final']=final;r['latency']=time.monotonic()-start;r['status']='FAIL'
  if final.get('state')=='completed':
   ops=final.get('result',{}).get('result',{}).get('operations',[])
   generated=next((o.get('result',{}) for o in ops if o.get('tool')=='document_create'),{})
   rid=generated.get('document_id');r['generated']=generated
   if rid:
    raw=urllib.request.urlopen('http://127.0.0.1:8088/documents/'+rid+'/original').read();name=case+('.docx' if case=='AT-15' else '.xlsx');(OUT/'artifacts'/name).write_bytes(raw);r['sha256']=hashlib.sha256(raw).hexdigest()
    if case=='AT-15':
     text='\n'.join(p.text for p in Document(BytesIO(raw)).paragraphs);r['decoded']=text;r['status']='PASS' if 'Mira' in text and '7' in text and 'prohibit' in text.lower() else 'FAIL'
    else:
     book=load_workbook(BytesIO(raw),data_only=False);rows=list(book.active.values);r['decoded']=rows;r['status']='PASS' if book.sheetnames==['Summary'] and rows[0]==('Group','Count') and set(rows[1:])=={('ALR',3),('ADR',1)} else 'FAIL'
 save('agent/'+case+'.json',r);results.append(r);print(case,r.get('status'),flush=True)
save('artifacts/source-bound-acceptance.json',results)
