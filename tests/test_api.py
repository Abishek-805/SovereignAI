from dataclasses import replace
from fastapi.testclient import TestClient
from backend.app import create_app, model_error_response
from tests.test_service import service

def client_for(service): return TestClient(create_app(service),base_url='http://127.0.0.1:8088')

def test_docker_start_requires_local_origin(service, monkeypatch):
    from backend import desktop
    calls = []
    monkeypatch.setattr(desktop, 'start_docker', lambda: calls.append(True) or {'status': 'starting'})
    with client_for(service) as client:
        assert client.post('/workbench/docker/start', headers={'Origin': 'https://example.com'}).status_code == 403
        assert not calls
        assert client.post('/workbench/docker/start', headers={'Origin': 'http://127.0.0.1:8088'}).json()['status'] == 'starting'
        assert calls == [True]


def test_model_load_requires_explicit_capability_and_local_origin(service, monkeypatch):
    calls = []
    monkeypatch.setattr(service.registry, 'acquire_lease', lambda capability: calls.append(capability))
    with client_for(service) as client:
        assert client.post('/workbench/model/load', json={'capability': 'coding'}).status_code == 422
        assert client.post('/workbench/model/load', json={'capability': 'text'},
                           headers={'Origin': 'https://example.com'}).status_code == 403
        response = client.post('/workbench/model/load', json={'capability': 'vision'},
                               headers={'Origin': 'http://127.0.0.1:8088'})
        assert response.status_code == 200, response.text
        assert calls == ['vision']

def test_context_overflow_has_actionable_message():
    import json
    content=json.dumps({'error':{'type':'exceed_context_size_error',
                                  'message':'request (5256 tokens) exceeds the available context size (4096 tokens)'}}).encode()
    response=model_error_response(content,400,'application/json')
    assert response.status_code==400
    assert 'Ask documents' in json.loads(response.body)['error']['message']
    untouched=model_error_response(b'{"error":{"message":"unrelated"}}',400,'application/json')
    assert untouched.body==b'{"error":{"message":"unrelated"}}'


def test_workbench_info_reports_unavailable_sandbox_without_claiming_execution(service):
    with client_for(service) as client:
        response=client.get('/workbench/info')
        assert response.status_code==200
        body=response.json()
        assert body['sandbox']['ready'] is False
        assert body['sandbox']['policy']['network']=='none'
        assert body['host']=='127.0.0.1'
        assert body['network_proof']!='air-gapped'
        assert next(tool for tool in body['tools'] if tool['name']=='Coding workspace / Docker')['available'] is False


def test_workbench_info_exposes_actual_capability_routes_and_model_files(service,tmp_path):
    from router.model_registry import ModelSpec
    text=ModelSpec('text','sovereign-text','text.gguf')
    vision=ModelSpec('vision','sovereign-vision','vision.gguf',modalities=('text','image'))
    (tmp_path/'text.gguf').write_bytes(b'model fixture')
    service.registry.specs={'text':text,'vision':vision}
    service.registry._model_paths=lambda spec:[tmp_path/spec.model_file]
    with client_for(service) as client:
        info=client.get('/workbench/info').json()
    assert info['routing']['mode']=='automatic'
    assert info['routing']['routes']['code']=={'capability':'text','model':'sovereign-text'}
    assert info['routing']['routes']['vision']=={'capability':'vision','model':'sovereign-vision'}
    assert {model['capability']:model['assets_present'] for model in info['models']}=={'text':True,'vision':False}


def test_coding_greeting_does_not_load_model_or_sandbox(service, monkeypatch):
    def unexpected(*args, **kwargs):
        raise AssertionError('Greeting must not read project files or load a runtime')
    monkeypatch.setattr(service, '_verified_coding_sandbox', unexpected)
    monkeypatch.setattr(service.registry, 'acquire_lease', unexpected)
    monkeypatch.setattr(service.coding, 'run_project', unexpected)
    for greeting in ('Hi!', 'hay'):
        result=service.run_coding_project_task('unused', '', greeting)
        assert result['state']=='answered'
        assert result['routing']['model'] is None


def test_artifact_catalog_withholds_download_when_hash_fails(service):
    import json
    task_id='a'*32
    directory=service.settings.data_dir.parent/'outputs'/task_id
    directory.mkdir(parents=True)
    (directory/'note.docx').write_bytes(b'changed')
    (directory/'manifest.json').write_text(json.dumps({
        'task_id':task_id,'created_at':'2026-09-26T00:00:00+00:00',
        'files':[{'name':'note.docx','sha256':'0'*64}]}))
    with client_for(service) as client:
        entry=client.get('/workbench/artifacts').json()[0]
        assert entry['validated'] is False and entry['url'] is None

def test_delete_download_removes_only_selected_artifact(service):
    import hashlib, json
    task_id='b'*32
    directory=service.settings.data_dir.parent/'outputs'/task_id
    directory.mkdir(parents=True)
    files=[]
    for name in ('first.docx','second.xlsx'):
        content=name.encode()
        (directory/name).write_bytes(content)
        files.append({'name':name,'sha256':hashlib.sha256(content).hexdigest()})
    (directory/'manifest.json').write_text(json.dumps({'task_id':task_id,'files':files}))
    with client_for(service) as client:
        assert client.delete(f'/workbench/artifacts/{task_id}/first.docx').json()=={'deleted':True}
        assert not (directory/'first.docx').exists()
        assert (directory/'second.xlsx').exists()
        assert [item['name'] for item in client.get('/workbench/artifacts').json()]==['second.xlsx']
        assert client.delete(f'/workbench/artifacts/{task_id}/first.docx').status_code!=200


def test_document_job_forwards_followup_context(service, monkeypatch):
    calls=[]
    monkeypatch.setattr(service,'ask',lambda question,document_ids,history: calls.append((question,document_ids,history)) or {'status':'answered','answer':'ok','sources':[]})
    with client_for(service) as client:
        job=client.post('/documents/jobs',json={'kind':'ask','question':'Compare it','document_ids':[],'history':['Summarize report A']}).json()
        import time
        for _ in range(40):
            state=client.get('/coding/jobs/'+job['job_id']).json()
            if state['state']!='running':break
            time.sleep(.025)
    assert state['state']=='completed'
    assert calls==[('Compare it',[],['Summarize report A'])]


def test_code_job_uses_project_task_without_requiring_open_file(service, monkeypatch):
    workspace_id=service.coding.create('Project edit')['workspace_id']
    calls=[]
    monkeypatch.setattr(service,'run_coding_project_task',lambda wid,target,instruction,job:
                        calls.append((wid,target,instruction)) or {'state':'completed','changes':[]})
    with client_for(service) as client:
        job=client.post(f'/coding/workspaces/{workspace_id}/jobs',json={
            'kind':'edit','target':'','instruction':'Create a new addition.py file'}).json()
        import time
        for _ in range(40):
            state=client.get('/coding/jobs/'+job['job_id']).json()
            if state['state']!='running':break
            time.sleep(.025)
    assert state['state']=='completed'
    assert calls==[(workspace_id,'','Create a new addition.py file')]

def test_ask_and_upload(service,tmp_path):
    with client_for(service) as client:
        assert client.post('/ask',json={'question':' '}).status_code==422
        response=client.post('/documents/import',files={'file':('manual.txt',b'P-101 limit is 7.1 mm/s.','text/plain')})
        assert response.status_code==200,response.text
        result=client.post('/ask',json={'question':'limit?'}).json()
        assert result['status']=='answered'
        assert result['sources'][0]['display_name']=='manual.txt'
        assert client.get('/documents').json()[0]['chunk_count']==1


def test_document_greeting_does_not_search_unrelated_files(service, monkeypatch):
    monkeypatch.setattr(service.embedder, 'encode', lambda *_: (_ for _ in ()).throw(AssertionError('greeting searched documents')))
    with client_for(service) as client:
        for greeting in ('hi', 'hay'):
            result=client.post('/ask',json={'question':greeting}).json()
            assert result['status']=='greeting'
            assert result['sources']==[]
            assert 'Ask me about your selected files' in result['answer']


def test_document_manager_rename_and_remove(service):
    with client_for(service) as client:
        response = client.post('/documents/import', files={'file': ('manage.txt', b'Local document management.')})
        assert response.status_code == 200, response.text
        document_id = response.json()['document_id']
        renamed = client.patch(f'/documents/{document_id}', json={'display_name': 'renamed.txt'})
        assert renamed.status_code == 200, renamed.text
        assert renamed.json()['display_name'] == 'renamed.txt'
        assert client.get('/documents').json()[0]['display_name'] == 'renamed.txt'
        removed = client.delete(f'/documents/{document_id}')
        assert removed.status_code == 200, removed.text
        assert client.get('/documents').json() == []


def test_coding_workspace_api_is_bounded(service):
    from router.sandbox import SandboxResult
    class Model:
        def complete_code(self, messages, max_tokens=1024):
            return {'code': 'def add(a, b):\n    return a + b\n'}
    class Sandbox:
        def _ready(self): pass
        def execute(self, code, input_files):
            assert input_files['solution.py'].endswith(b'return a + b\n')
            return SandboxResult(0, '', 'Ran 1 test', True, {'result.txt': b'ok'})
    service.model = Model()
    service._verified_coding_sandbox = lambda: Sandbox()
    with client_for(service) as client:
        created = client.post('/coding/workspaces',json={'name':'Unit task'})
        assert created.status_code == 200
        workspace_id = created.json()['workspace_id']
        for name, content in [('solution.py','def add(a, b):\n    return 0\n'),
                              ('test_solution.py','import unittest\n')]:
            saved = client.put(f'/coding/workspaces/{workspace_id}/files/{name}',json={'content':content})
            assert saved.status_code == 200, saved.text
        assert len(client.get(f'/coding/workspaces/{workspace_id}').json()['files']) == 2
        assert client.get(f'/coding/workspaces/{workspace_id}/files/solution.py').json()['content'].endswith('return 0\n')
        assert client.put(f'/coding/workspaces/{workspace_id}/files/%2e%2e%2fsecret.py',
                          json={'content':'x'}).status_code in {400,404,405}
        run = client.post(f'/coding/workspaces/{workspace_id}/tasks',
                          json={'target':'solution.py','instruction':'Fix add'})
        assert run.status_code == 200, run.text
        result = run.json()
        assert result['state'] == 'completed' and '+    return a + b' in result['diff']
        assert client.get(f"/coding/workspaces/{workspace_id}/tasks/{result['task_id']}").json()['checks']['tests_passed']
        assert client.get(result['output_files'][0]['url']).content == b'ok'
        artifact_name=result['output_files'][0]['name']
        assert client.delete(f"/workbench/artifacts/{result['task_id']}/{artifact_name}").json()=={'deleted':True}
        assert all(item['name']!=artifact_name or item['task_id']!=result['task_id']
                   for item in client.get('/workbench/artifacts').json())
        undone=client.post(f"/coding/workspaces/{workspace_id}/tasks/{result['task_id']}/undo")
        assert undone.status_code==200 and undone.json()['state']=='undone'
        assert client.get(f'/coding/workspaces/{workspace_id}/files/solution.py').json()['content'].endswith('return 0\n')


def test_docx_preview_and_import(service):
    from io import BytesIO
    from docx import Document
    document=Document(); document.add_paragraph('Unit 3 equation: x squared plus one.')
    stream=BytesIO(); document.save(stream); data=stream.getvalue()
    with client_for(service) as client:
        preview=client.post('/documents/preview',files={'file':('problems.docx',data)})
        assert preview.status_code==200,preview.text
        assert 'x squared plus one' in preview.json()['text']
        indexed=client.post('/documents/import',files={'file':('problems.docx',data)})
        assert indexed.status_code==200,indexed.text
        assert client.get('/documents').json()[0]['display_name']=='problems.docx'
        chunk=service.store.active_chunks()[0][0]
        original=client.get('/sources/'+chunk.chunk_id+'/original')
        assert original.status_code==200 and original.content==data


def test_server_tool_contract_and_safe_search(service):
    with client_for(service) as client:
        listed=client.get('/tools')
        assert listed.status_code==200
        assert any(tool['tool']=='search_documents' for tool in listed.json())
        imported=client.post('/documents/import',files={'file':('sop.txt',b'Pump A-17 limit is 7.0 mm/s.')}).json()
        result=client.post('/tools',json={'tool':'search_documents','params':{'question':'A-17 limit'}})
        assert result.status_code==200
        assert '7.0 mm/s' in result.json()['plain_text_response']
        assert 'sop.txt' in result.json()['plain_text_response']
        named=client.post('/tools',json={'tool':'search_documents','params':{'question':'A-17 limit','document_ids':['sop.txt']}})
        assert named.status_code==200
        assert '7.0 mm/s' in named.json()['plain_text_response']
        rejected=client.post('/tools',json={'tool':'read_file','params':{'path':'C:/secret'}})
        assert rejected.status_code==200
        assert 'error' in rejected.json()
        streamed=client.post('/tools',json={'tool':'search_documents','params':{'question':'A-17 limit'},'stream':True})
        assert streamed.status_code==200
        assert '"done":true' in streamed.text

def test_hostile_filename(service,tmp_path):
    with client_for(service) as client:
        response=client.post('/documents/import',files={'file':('../../outside.txt',b'hello','text/plain')})
        assert response.status_code==400
    assert not (tmp_path/'outside.txt').exists()

def test_origins_hosts_and_busy(service):
    with client_for(service) as client:
        assert client.post('/ask',json={'question':'hi'},headers={'Origin':'https://evil.example'}).status_code==403
        assert client.get('/documents',headers={'Host':'evil.example'}).status_code==400
        service.ask_lock.acquire()
        try:
            response=client.post('/ask',json={'question':'hi'})
            assert response.status_code==409
            assert response.json()['code']=='busy'
        finally: service.ask_lock.release()

def test_upload_limits_and_extension(service):
    service.settings=replace(service.settings,max_file_bytes=20)
    with client_for(service) as client:
        assert client.post('/documents/import',files={'file':('a.txt',b'a'*21)}).status_code==413
        assert client.post('/documents/import',files={'file':('a.exe',b'a')}).status_code==400
        assert client.post('/documents/import',content=b'a'*100000,headers={'Content-Type':'application/octet-stream'}).status_code==413

def test_unavailable_not_insufficient(service):
    from backend.contracts import WorkbenchError
    def unavailable(*args,**kwargs): raise WorkbenchError('model_unavailable','unavailable')
    service.ask=unavailable
    with client_for(service) as client:
        assert client.post('/ask',json={'question':'hi'}).status_code==503

def test_repeated_upload_is_unchanged(service):
    with client_for(service) as client:
        first=client.post('/documents/import',files={'file':('same.txt',b'limit 7.1')}).json()
        second=client.post('/documents/import',files={'file':('same.txt',b'limit 7.1')}).json()
        assert second['status']=='unchanged'
        assert first['document_id']==second['document_id']
        different=client.post('/documents/import',files={'file':('same.txt',b'other content')}).json()
        assert different['document_id']!=first['document_id']

def test_api_has_no_external_docs_assets(service):
    with client_for(service) as client:
        assert client.get('/docs').status_code==404


def test_llama_ui_is_served_and_model_routes_are_local(service, monkeypatch):
    import httpx
    calls=[]
    class FakeClient:
        async def __aenter__(self): return self
        async def __aexit__(self,*args): pass
        async def request(self,method,url,**kwargs):
            calls.append((method,url,kwargs))
            return httpx.Response(200,json={'ok':True})
    monkeypatch.setattr('backend.app.httpx.AsyncClient',lambda **kwargs: FakeClient())
    with client_for(service) as client:
        page=client.get('/')
        assert page.status_code==200
        assert '/_app/' in page.text
        assert client.get('/props').json()=={'ok':True}
        assert client.get('/v1/models').json()=={'ok':True}
        assert client.post('/v1/chat/completions',json={'messages':[]}).json()=={'ok':True}
        assert client.get('/v1/not-allowed').status_code==404
    assert [url for _,url,_ in calls]==['http://127.0.0.1:8087/props','http://127.0.0.1:8087/v1/models','http://127.0.0.1:8087/v1/chat/completions']


def test_sleeping_model_keeps_last_known_ui_metadata(service, monkeypatch):
    import httpx
    calls=0
    class FakeClient:
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        async def request(self,method,url,**kwargs):
            nonlocal calls
            calls+=1
            return httpx.Response(200,json={'model_alias':'sovereign-text'}) if calls==1 else httpx.Response(503,json={'error':'sleeping'})
    monkeypatch.setattr('backend.app.httpx.AsyncClient',lambda **kwargs:FakeClient())
    with client_for(service) as client:
        assert client.get('/props').json()=={'model_alias':'sovereign-text'}
        assert client.get('/props').json()=={'model_alias':'sovereign-text'}


def test_document_job_can_be_rejoined_after_navigation(service, monkeypatch):
    import time
    monkeypatch.setattr(service,'ask',lambda question,document_ids:{'status':'answered','answer':question,'sources':[]})
    with client_for(service) as client:
        started=client.post('/documents/jobs',json={'kind':'ask','question':'Where is the source?','document_ids':[]})
        assert started.status_code==200
        job_id=started.json()['job_id']
        for _ in range(40):
            state=client.get('/coding/jobs/'+job_id).json()
            if state['state']!='running':break
            time.sleep(.025)
        assert state['state']=='completed'
        assert state['result']['answer']=='Where is the source?'
        assert client.get('/coding/jobs/'+job_id).json()['result']==state['result']


def test_general_chat_does_not_overlap_document_or_model_switch(service):
    with client_for(service) as client:
        service.ask_lock.acquire()
        try:
            response=client.post('/v1/chat/completions',json={'messages':[]})
            assert response.status_code==409
            assert response.json()['code']=='busy'
        finally:
            service.ask_lock.release()


def test_maintenance_draft_and_download(service,tmp_path):
    with client_for(service) as client:
        report=client.post('/documents/import',files={'file':('report.txt',b'Pump P-101 vibration measured 8.2 mm/s.')}).json()
        sop=client.post('/documents/import',files={'file':('sop.txt',b'Investigate Pump P-101 vibration above 7.1 mm/s.')}).json()
        response=client.post('/workflows/maintenance-draft',json={'document_ids':[report['document_id'],sop['document_id']]})
        assert response.status_code==200,response.text
        result=response.json()
        assert result['status']=='completed'
        assert result['findings'][0]['status']=='exceeds investigation threshold'
        trace=client.get('/tasks/'+result['task_id'])
        assert trace.status_code==200
        assert trace.json()['state']=='completed'
        assert [step['name'] for step in trace.json()['steps']]==['select_evidence','compare_measurements','write_deliverables']
        assert all(trace.json()['checks'].values())
        assert client.get('/tasks/bad').status_code==404
        word=client.get(result['downloads']['word'])
        excel=client.get(result['downloads']['excel'])
        slides=client.get(result['downloads']['slides'])
        assert word.status_code==200 and word.content[:2]==b'PK'
        assert excel.status_code==200 and excel.content[:2]==b'PK'
        assert slides.status_code==200 and slides.content[:2]==b'PK'
        artifacts=client.get('/workbench/artifacts').json()
        assert {item['kind'] for item in artifacts if item['task_id']==result['task_id']}=={'word','excel','slides'}
        assert all(item['validated'] for item in artifacts if item['task_id']==result['task_id'])
        assert client.get('/artifacts/bad/file.docx').status_code==404


def test_agent_routes_and_checks_maintenance_deliverables(service):
    with client_for(service) as client:
        report=client.post('/documents/import',files={'file':('report.txt',b'Pump P-101 vibration measured 8.2 mm/s.')}).json()
        sop=client.post('/documents/import',files={'file':('sop.txt',b'Investigate Pump P-101 vibration above 7.1 mm/s.')}).json()
        response=client.post('/agent/tasks',json={'goal':'Compare the inspection report with the SOP and draft a maintenance note',
                                                  'document_ids':[report['document_id'],sop['document_id']]})
        assert response.status_code==200,response.text
        result=response.json()
        assert result['workflow']=='maintenance_draft'
        assert result['selected_model'] is None
        assert result['task_id']!=result['child_task_id']
        assert client.get('/tasks/'+result['task_id']).json()['state']=='completed'
        assert client.get(result['result']['downloads']['slides']).status_code==200


def test_agent_rejects_unapproved_goal_and_records_failure(service):
    with client_for(service) as client:
        response=client.post('/agent/tasks',json={'goal':'Control the plant machinery'})
        assert response.status_code==400
    tasks=list((service.settings.data_dir/'tasks').glob('*.json'))
    assert len(tasks)==1
    assert service.tasks.read(tasks[0].stem)['error_code']=='unsupported_goal'


def test_agent_calculation_shows_steps_and_rejects_unsafe_expression(service):
    with client_for(service) as client:
        response=client.post('/agent/tasks',json={'goal':'Calculate: (8.2 - 7.1) * 2'})
        assert response.status_code==200,response.text
        result=response.json()
        assert result['result']['rounded']==2.2
        assert result['selected_model'] is None
        assert result['result']['steps']==['8.2 - 7.1 = 1.1','1.1 × 2 = 2.2']
        assert client.get('/tasks/'+result['task_id']).json()['state']=='completed'
        assert client.post('/calculate',json={'expression':'__import__("os")'}).status_code==400


def test_agent_document_answer_is_grounded_and_persisted(service):
    with client_for(service) as client:
        doc=client.post('/documents/import',files={'file':('limit.txt',b'Pump P-101 limit is 7.1 mm/s.')}).json()
        response=client.post('/agent/tasks',json={'goal':'What is the limit for Pump P-101?',
                                                   'document_ids':[doc['document_id']]})
        assert response.status_code==200,response.text
        body=response.json()
        assert body['capability']=='DOCUMENT_QA' and body['checks']['grounded']
        assert body['route']['model']=='sovereign-text'
        assert body['resource']['retrieved_evidence']>=1
        assert body['resource']['max_steps']==8
        saved=client.get('/tasks/'+body['task_id']).json()
        assert saved['agent_state']=='completed' and saved['agent_result']['capability']=='DOCUMENT_QA'
        assert saved['agent_history'][-1]['state']=='completed'


def test_agent_workspace_uses_existing_docker_boundary(service):
    from router.sandbox import SandboxResult
    class Model:
        def complete_code(self,messages,max_tokens=1024):
            return {'code':'def add(a, b):\n    return a + b\n'}
    class Sandbox:
        image_id='sha256:'+'a'*64
        def _ready(self): pass
        def execute(self,code,input_files):
            assert input_files['solution.py'].endswith(b'return a + b\n')
            return SandboxResult(0,'','Ran 1 test',True,{'result.txt':b'ok'})
    service.model=Model(); service._verified_coding_sandbox=lambda: Sandbox()
    with client_for(service) as client:
        workspace=client.post('/coding/workspaces',json={'name':'Agent code'}).json()['workspace_id']
        for name,content in [('solution.py','def add(a,b):\n    return 0\n'),('test_solution.py','import unittest\n')]:
            assert client.put(f'/coding/workspaces/{workspace}/files/{name}',json={'content':content}).status_code==200
        response=client.post('/agent/tasks',json={'goal':'Fix add','workspace_id':workspace,'target':'solution.py'})
        assert response.status_code==200,response.text
        body=response.json()
        assert body['capability']=='CODING' and body['checks']['child_checks_passed']
        assert body['result']['sandbox']['backend']=='docker'
        assert body['artifacts'][0]['name']=='result.txt'


def test_agent_failed_tool_records_failure(service):
    with client_for(service) as client:
        response=client.post('/agent/tasks',json={'goal':'Fix parser','workspace_id':'bad','target':'parser.py'})
        assert response.status_code==400
        records=list((service.settings.data_dir/'tasks').glob('*.json'))
        assert len(records)==1
        saved=service.tasks.read(records[0].stem)
        assert saved['agent_state']=='failed' and saved['error_code']=='sandbox_unavailable'


def test_agent_needs_input_exposes_saved_state(service):
    with client_for(service) as client:
        response=client.post('/agent/tasks',json={'goal':'Draft a maintenance note',
                                                   'document_ids':['one-document']})
        assert response.status_code==400
        body=response.json()
        assert body['code']=='needs_input'
        saved=client.get('/tasks/'+body['task_id']).json()
        assert saved['state']=='needs_input' and saved['agent_state']=='needs_input'


def test_agent_image_uses_vision_lease(service,monkeypatch):
    from io import BytesIO
    from PIL import Image
    leases=[]
    service.registry.acquire_lease=lambda capability: leases.append(capability)
    monkeypatch.setattr('backend.service.ask_vision',lambda image,question:
                        {'status':'answered','answer':'A pump label'})
    image=BytesIO(); Image.new('RGB',(20,20),'white').save(image,format='PNG')
    with client_for(service) as client:
        response=client.post('/agent/vision',files={'file':('label.png',image.getvalue(),'image/png')},
                             data={'question':'What is visible?'})
        assert response.status_code==200,response.text
        body=response.json()
        assert body['capability']=='VISION' and body['result']['answer']=='A pump label'
        assert leases==['vision']
        assert client.get('/tasks/'+body['task_id']).json()['agent_state']=='completed'


def test_failed_maintenance_draft_keeps_failed_trace(service,tmp_path):
    with client_for(service) as client:
        document=client.post('/documents/import',files={'file':('empty-pattern.txt',b'No equipment measurement is recorded here.')}).json()
        before=set((service.settings.data_dir/'tasks').glob('*.json')) if (service.settings.data_dir/'tasks').exists() else set()
        response=client.post('/workflows/maintenance-draft',json={'document_ids':[document['document_id']]})
        assert response.status_code==400
        created=set((service.settings.data_dir/'tasks').glob('*.json'))-before
        assert len(created)==1
        trace=client.get('/tasks/'+created.pop().stem).json()
        assert trace['state']=='failed'
        assert trace['error_code']=='missing_evidence'
        assert 'finished_at' in trace


def test_historical_source_and_original_remain_inspectable(service,tmp_path):
    source=tmp_path/'report.txt';source.write_text('P-101 vibration measured 8.2 mm/s.',encoding='utf-8')
    first=service.import_file(source)
    chunk=service.store.active_chunks()[0][0]
    source.write_text('P-101 vibration measured 8.5 mm/s.',encoding='utf-8')
    service.import_file(source)
    with client_for(service) as client:
        old=client.get('/sources/'+chunk.chunk_id)
        assert old.status_code==200
        assert old.json()['version_hash']==first['active_hash']
        assert old.json()['text']=='P-101 vibration measured 8.2 mm/s.'
        original=client.get('/sources/'+chunk.chunk_id+'/original')
        assert original.status_code==200
        assert original.content==b'P-101 vibration measured 8.2 mm/s.'


def test_vision_upload_validates_file_and_returns_answer(service):
    from io import BytesIO
    from PIL import Image
    image=BytesIO()
    Image.new('RGB',(60,40),'white').save(image,format='PNG')
    service.ask_vision=lambda path,question: {'status':'answered','answer':'P-101'}
    with client_for(service) as client:
        response=client.post('/vision/ask',data={'question':'Read the tag'},files={'file':('tag.png',image.getvalue(),'image/png')})
        assert response.status_code==200,response.text
        assert response.json()['answer']=='P-101'
        assert client.post('/vision/ask',data={'question':'Read'},files={'file':('tag.exe',b'junk')}).status_code==400


def test_document_content_and_generic_report(service):
    with client_for(service) as client:
        doc=client.post('/documents/import',files={'file':('manual.txt',b'P-101 limit is 7.1 mm/s.','text/plain')}).json()
        content=client.get('/documents/'+doc['document_id']+'/content')
        assert content.status_code==200
        assert '7.1' in content.json()['text'] and content.json()['methods']==['utf8_text']
        report=client.post('/documents/report',json={'question':'What is the limit?','document_ids':[doc['document_id']]})
        assert report.status_code==200,report.text
        assert client.get(report.json()['downloads']['word']).status_code==200
        wid=client.post('/coding/workspaces',json={'name':'Nested project'}).json()['workspace_id']
        assert client.put(f'/coding/workspaces/{wid}/files/src%2Fmain.js',json={'content':'console.log(42)'}).status_code==200
        assert client.get(f'/coding/workspaces/{wid}/files/src%2Fmain.js').json()['content']=='console.log(42)'


def test_auto_agent_dispatches_calculation_without_manual_fields(service):
    service.model.plan_task=lambda goal,docs,files,history: {'action':'calculate','target':'','expression':'18.5 * 24','response':'Calculate exactly'}
    with client_for(service) as client:
        response=client.post('/agent/auto',json={'goal':'Calculate 18.5 times 24'})
        assert response.status_code==200,response.text
        assert '444' in response.json()['answer']
        assert response.json()['status']=='completed'


def test_auto_agent_reads_selected_documents_even_if_planner_tries_to_guess(service):
    service.model.plan_task=lambda *args: {'action':'answer','target':'','expression':'','response':'Invented answer'}
    with client_for(service) as client:
        doc=client.post('/documents/import',files={'file':('manual.txt',b'P-101 limit is 7.1 mm/s.','text/plain')}).json()
        result=client.post('/agent/auto',json={'goal':'What is the limit?','document_ids':[doc['document_id']]}).json()
        assert result['plan']['action']=='search_documents'
        assert result['result']['sources'] and result['answer']!='Invented answer'


def test_auto_agent_reads_named_document_without_explicit_selection(service):
    service.model.plan_task=lambda *args: {'action':'answer','target':'','expression':'','response':'I cannot access that PDF'}
    with client_for(service) as client:
        doc=client.post('/documents/import',files={'file':('24ALR001_notes.txt',b'The report describes database architecture and machine learning.','text/plain')}).json()
        result=client.post('/agent/auto',json={'goal':'Tell me about the 24alr001 file'}).json()
        assert result['plan']['action']=='search_documents'
        assert result['result']['sources'][0]['document_id']==doc['document_id']
        assert result['answer']!='I cannot access that PDF'


def test_auto_agent_explanation_receives_source_and_does_not_edit(service):
    wid=service.coding.create('Explain source')['workspace_id']
    service.coding.write(wid,'src/main.js','console.log(42)')
    def plan(goal,docs,files,history):
        assert files[0]['source_excerpt']=='console.log(42)'
        return {'action':'answer','target':'','expression':'','response':'Prints 42.'}
    service.model.plan_task=plan
    result=service.run_auto_agent('Explain this source',workspace_id=wid)
    assert result['answer']=='Prints 42.'
    assert service.coding.read(wid,'src/main.js')['content']=='console.log(42)'


def test_auto_agent_routes_explicit_error_repair_to_code_tool(service):
    wid=service.coding.create('Broken code')['workspace_id']
    service.coding.write(wid,'broken.py','print(sales)\n')
    service.model.plan_task=lambda *args: {'action':'answer','target':'','expression':'','response':'I fixed it'}
    service._verified_coding_sandbox=lambda: object()
    calls=[]
    def repair(workspace_id,target,instruction,job=None):
        calls.append((workspace_id,target,instruction))
        return {'state':'completed','target':target,'validation':'runtime_check','checks':{'runtime_passed':True}}
    service.run_coding_project_task=repair
    result=service.run_auto_agent('Solve the errors in the code',workspace_id=wid)
    assert result['plan']['action']=='edit_code'
    assert calls and calls[0][:2]==(wid,'broken.py')


def test_spreadsheet_original_is_retrievable(service):
    from openpyxl import Workbook
    from io import BytesIO
    book=Workbook();book.active.append(['limit',7.1]);buffer=BytesIO();book.save(buffer)
    with client_for(service) as client:
        doc=client.post('/documents/import',files={'file':('limits.xlsx',buffer.getvalue(),'application/octet-stream')}).json()
        preview=client.get('/documents/'+doc['document_id']+'/content').json()
        response=client.get(preview['original_url'])
        assert response.status_code==200
        assert response.content==buffer.getvalue()


def test_pdf_original_displays_inline(service):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    from io import BytesIO
    buffer=BytesIO();pdf=PdfWriter();page=pdf.add_blank_page(600,800)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 30 700 Td (P-101 limit is 7.1 mm/s.) Tj ET')
    page[NameObject('/Contents')]=pdf._add_object(stream);pdf.write(buffer)
    with client_for(service) as client:
        doc=client.post('/documents/import',files={'file':('manual.pdf',buffer.getvalue(),'application/pdf')}).json()
        preview=client.get('/documents/'+doc['document_id']+'/content').json()
        response=client.get(preview['original_url'])
        assert response.status_code==200 and response.headers['content-disposition'].startswith('inline')
