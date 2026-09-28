"""Planner-selected overview covers document scope rather than top-six search hits."""
import json

from backend.contracts import Chunk
from rag.answer import answer
from rag.overview import overview_passages, overview_documents
from tests.test_answer import FakeModel
from tests.test_service import service


def test_overview_uses_all_connected_documents_without_embedding_query(service, tmp_path):
    ids=[]
    for index in range(16):
        path=tmp_path/f'File_{index}.txt'
        path.write_text(f'Document topic: equipment group {index}.')
        ids.append(service.import_file(path)['document_id'])
    extra=tmp_path/'Unconnected.txt'
    extra.write_text('This document is outside the connected scope.')
    extra_id=service.import_file(extra)['document_id']
    service.model.plan_task=lambda *_:{'action':'search_documents','document_scope':'overview',
                                      'response':'Synthesize selected documents'}
    service.model.output={'status':'answered','answer':'The supplied excerpts describe equipment groups [S1].'}
    calls_before=service.embedder.calls
    result=service.ask('Give an overview of all my connected documents',ids)
    assert result['status']=='answered'
    assert {source['document_id'] for source in result['sources']}==set(ids)
    assert extra_id not in {source['document_id'] for source in result['sources']}
    assert service.embedder.calls==calls_before
    assert result['coverage']['requested_documents']==16
    assert result['coverage']['covered_documents']==16
    assert result['coverage']['all_documents_represented'] is True
    prompt=service.model.calls[-1]
    assert 'You create the summary; it need not already exist' in prompt[0]['content']
    evidence=json.loads(prompt[-1]['content'].split('\nAnswer the current question.')[0])
    assert evidence['evidence_coverage']==result['coverage']


def test_balancing_offers_each_document_before_additional_pages():
    chunks=[(Chunk(f'{doc}-{page}',doc,'version',doc+'.pdf','Evidence',page,1,1),None)
            for doc in ('A','B','C') for page in range(1,6)]
    selected=overview_passages(chunks)
    assert [chunk.document_id for chunk in selected[:3]]==['A','B','C']
    assert len(selected)==9
    assert [chunk.page for chunk in selected[:3]]==[1,1,1]


def test_context_limited_overview_identifies_omitted_documents_and_passages():
    active=[(Chunk(str(index),str(index),'version',f'Doc{index}.txt',
                   ('First supported observation.' if index==0 else 'Long supporting text '*500),1,1,1),None)
            for index in range(4)]
    model=FakeModel()
    model.output={'status':'answered','answer':'The first excerpt records an observation [S1].'}
    result=answer('Summarize the collection',overview_passages(active),model,
                  context=1800,output_tokens=200,overview_documents=overview_documents(active))
    assert len(result['sources'])==1
    assert result['coverage']['covered_documents']==1
    assert result['coverage']['requested_documents']==4
    assert result['coverage']['all_documents_represented'] is False
    assert result['coverage']['all_indexed_passages_included'] is False
    assert [doc['included_passages'] for doc in result['coverage']['documents']]==[1,0,0,0]
    assert model.count_messages(model.calls[-1])+200+64<=1800


def test_agent_propagates_model_scope_into_document_workflows(service,monkeypatch):
    monkeypatch.setattr(service,'documents',lambda:[{'document_id':'selected','display_name':'Selected.txt'}])
    calls=[]
    def workflow(question,ids,**kwargs):
        calls.append(kwargs)
        return {'status':'answered','answer':'A supplied observation [S1].','sources':[]}
    for action in ('search_documents','create_report'):
        service.model.plan_task=lambda *_,chosen=action:{'action':chosen,'document_scope':'overview',
                                                       'response':'Synthesize connected documents','target':'','expression':''}
        monkeypatch.setattr(service,'ask' if action=='search_documents' else 'create_document_report',workflow)
        service.run_auto_agent('Summarize all connected documents',['selected'])
    assert all(call['document_scope']=='overview' for call in calls)


def test_overview_report_uses_model_synthesis_with_coverage(service,tmp_path,monkeypatch):
    from docx import Document
    path=tmp_path/'File.txt';path.write_text('Supported observation: assembly is documented.')
    document_id=service.import_file(path)['document_id']
    service.model.output={'status':'answered','answer':'Assembly is documented in the supplied excerpt [S1].'}
    result=service.create_document_report('Generate a report',[document_id],document_scope='overview')
    assert result['answer']==service.model.output['answer']
    assert result['coverage']['covered_documents']==1
    assert result['downloads']['word'].endswith('/word')
    assert service.model.calls
    report=Document(service.settings.data_dir.parent/'outputs'/result['task_id']/'Document report.docx')
    text='\n'.join(paragraph.text for paragraph in report.paragraphs)
    assert 'Evidence coverage' in text
    assert '1 of 1 connected documents are represented by 1 of 1 indexed passages' in text
    assert 'does not verify that all original file content' in text


def test_word_coverage_discloses_omitted_and_partial_documents_even_if_answer_does_not(service):
    from docx import Document
    from workflows.document_report import publish_report
    result={'status':'answered','answer':'A supported observation [S1].',
            'sources':[{'label':'S1','display_name':'Read.txt','page':None,'text':'A supported observation.'}],
            'coverage':{'mode':'overview','documents':[
                {'display_name':'Read.txt','indexed_passages':4,'included_passages':1},
                {'display_name':'Omitted.txt','indexed_passages':2,'included_passages':0},
            ]}}
    publish_report(service.settings,'coverage-test','Summarize connected documents',result)
    report=Document(service.settings.data_dir.parent/'outputs'/'coverage-test'/'Document report.docx')
    text='\n'.join(paragraph.text for paragraph in report.paragraphs)
    assert '1 of 2 connected documents are represented by 1 of 6 indexed passages' in text
    assert 'Documents without included evidence: Omitted.txt' in text
    assert 'Read.txt (1 of 4 passages)' in text
