import pytest
from pathlib import Path
from rag.deliverables import write_deliverables, _sanitize_excel_cell
from backend.settings import Settings
from backend.contracts import WorkbenchError

@pytest.fixture
def deliverable_data():
    return {
        "task_id": "test_task_123",
        "title": "Maintenance Approval Note",
        "document_date": "2026-09-23T12:00:00Z",
        "source_versions": [{"document_id": "doc1", "version_hash": "hash1"}],
        "findings": [
            {
                "observation": "Limit is 7.1",
                "requirement": "Must be < 8.0",
                "status": "supported",
                "source_ids": ["doc1"],
                "calc_id": "calc1"
            }
        ],
        "calculations": [
            {
                "id": "calc1",
                "expression": "8.0 - 7.1",
                "result": 0.9,
                "rounded": 0.9
            }
        ],
        "requested_action": "Approve maintenance.",
        "open_questions": ["Is this the latest limit?"],
        "references": [{"id": "S1", "description": "SOP Manual"}]
    }

def test_generate_deliverables(deliverable_data, tmp_path):
    import docx
    import openpyxl
    from pptx import Presentation

    settings = Settings(data_dir=tmp_path / 'data')
    result = write_deliverables(deliverable_data["task_id"], deliverable_data, settings)
    
    word_path = Path(result["word_path"])
    excel_path = Path(result["excel_path"])
    slide_path = Path(result["slide_path"])
    
    assert word_path.exists()
    assert word_path.name == "Maintenance Approval Note.docx"
    assert word_path.stat().st_size > 1000, "DOCX file must not be empty or placeholder"
    
    # Programmatically parse and validate DOCX
    doc = docx.Document(str(word_path))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith('Heading')]
    assert "Requested Action" in headings
    assert "Findings" in headings
    assert len(doc.tables) >= 1
    table = doc.tables[0]
    row_text = [cell.text for cell in table.rows[1].cells]
    assert "Limit is 7.1" in row_text[0]
    assert "doc1" in row_text[3]
    
    assert excel_path.exists()
    assert excel_path.name == "Maintenance Approval Note_evidence.xlsx"
    assert excel_path.stat().st_size > 1000, "XLSX file must not be empty or placeholder"
    
    # Programmatically parse and validate XLSX
    wb = openpyxl.load_workbook(str(excel_path))
    assert "Findings" in wb.sheetnames
    assert "Calculations" in wb.sheetnames
    assert "Sources" in wb.sheetnames
    assert "Unresolved Issues" in wb.sheetnames
    
    findings_row = [cell.value for cell in wb["Findings"][2]]
    assert findings_row[0] == "Limit is 7.1"
    assert findings_row[3] == "doc1"
    
    calc_row = [cell.value for cell in wb["Calculations"][2]]
    assert calc_row[0] == "calc1"
    assert calc_row[2] == 0.9
    assert slide_path.exists() and slide_path.stat().st_size > 1000
    slides=Presentation(slide_path)
    slide_text=' '.join(shape.text for slide in slides.slides for shape in slide.shapes if shape.has_text_frame)
    assert len(slides.slides)>=3
    assert 'S1' in slide_text and 'Is this the latest limit?' in slide_text

def test_missing_schema_fields(deliverable_data, tmp_path):
    del deliverable_data["findings"]
    settings = Settings(data_dir=tmp_path / 'data')
    with pytest.raises(WorkbenchError) as exc:
        write_deliverables("test_task_123", deliverable_data, settings)
    assert "Missing required field" in str(exc.value)

def test_sanitize_excel():
    assert _sanitize_excel_cell("=CMD()") == "'=CMD()"
    assert _sanitize_excel_cell("+100") == "'+100"
    assert _sanitize_excel_cell("Normal string") == "Normal string"
    assert _sanitize_excel_cell(100) == 100


def test_task_id_cannot_escape_output_root(deliverable_data,tmp_path):
    with pytest.raises(WorkbenchError):
        write_deliverables('../outside',deliverable_data,Settings(data_dir=tmp_path/'data'))
    assert not (tmp_path/'outside').exists()


def test_failed_excel_generation_publishes_no_word_file(deliverable_data,tmp_path,monkeypatch):
    import rag.deliverables as deliverables
    def fail(*args): raise RuntimeError('simulated workbook failure')
    monkeypatch.setattr(deliverables,'generate_excel_evidence',fail)
    settings=Settings(data_dir=tmp_path/'data')
    with pytest.raises(WorkbenchError):
        write_deliverables(deliverable_data['task_id'],deliverable_data,settings)
    assert not (tmp_path/'outputs'/deliverable_data['task_id']).exists()
