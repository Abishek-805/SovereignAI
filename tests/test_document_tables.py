import threading
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from tests.test_service import service
from tests.test_api import client_for


def workbook(tmp_path):
    book=Workbook(); sheet=book.active; sheet.title='Assessment'
    sheet.append(['Training results']); sheet.merge_cells('A1:C1')
    sheet.append(['Student ID','Total (50M)','Result'])
    sheet.append(['24ALR001',36.5,'PASS']); sheet.append(['24ALR002',12,'FAIL'])
    sheet['A2'].font=Font(bold=True); sheet['A2'].fill=PatternFill('solid',fgColor='FF123456')
    path=tmp_path/'professional-training.xlsx'; book.save(path); return path


def test_upload_is_visible_and_previewable_before_embedding_finishes(service,tmp_path):
    path=workbook(tmp_path); entered=threading.Event(); release=threading.Event()
    original=service.embedder.encode
    def slow(*args,**kwargs):
        entered.set(); assert release.wait(5); return original(*args,**kwargs)
    service.embedder.encode=slow
    with client_for(service) as client:
        try:
            job=client.post('/documents/import?background=true',files={'file':(path.name,path.read_bytes())}).json()
            assert entered.wait(2)
            doc=client.get('/documents').json()[0]
            assert doc['indexing_status']=='indexing'
            assert client.get('/documents/'+doc['document_id']+'/original').content==path.read_bytes()
            preview=client.get('/documents/'+doc['document_id']+'/table').json()
            assert preview['rows'][2]['cells'][1]['value']==36.5
            assert preview['merges']==['A1:C1']
            assert preview['rows'][1]['cells'][0]['bold'] is True
        finally: release.set()


def test_student_id_does_not_exclude_workbook_when_it_names_a_pdf(service,tmp_path):
    from rag.retrieve import document_scope_for_question
    docs=[{'document_id':'pdf','display_name':'24ALR001_ABISHEK.pdf'},
          {'document_id':'xlsx','display_name':'professional-training.xlsx'}]
    assert document_scope_for_question(docs,'mark of 24ALR001 in assessment') is None
    assert document_scope_for_question(docs,'in professional-training.xlsx')==['xlsx']


def test_table_queries_use_complete_records_and_validate_columns(service,tmp_path):
    from rag.tables import load_tables, execute_query
    tables=load_tables(workbook(tmp_path))
    table=tables[0]
    assert table['columns']==['Student ID','Total (50M)','Result']
    result=execute_query(table,{'operation':'count','filters':[{'column':'Result','operator':'eq','value':'PASS'}],'columns':[]})
    assert result['value']==1 and result['scanned_rows']==2
    result=execute_query(table,{'operation':'select','filters':[{'column':'Student ID','operator':'eq','value':'24ALR001'}],'columns':['Total (50M)']})
    assert result['records']==[{'Total (50M)':36.5}]
    with pytest.raises(Exception):
        execute_query(table,{'operation':'count','filters':[{'column':'invented','operator':'eq','value':'PASS'}],'columns':[]})


def test_grounded_answer_uses_verified_aggregate_and_readable_units(service,tmp_path):
    from rag.answer import _unsupported_numbers
    doc=service.import_file(workbook(tmp_path))
    def plan(question,tables,history):
        return {'table':tables[0]['id'],'operation':'count','columns':[],
                'filters':[{'column':'Result','operator':'eq','value':'PASS'}]}
    service.model.plan_table_query=plan
    captured=[]
    service.model.complete=lambda messages,max_tokens: captured.extend(messages) or {'result':{'status':'answered','answer':'1 student passed. [S1]'}}
    result=service.ask('How many passed?',[doc['document_id']],force_documents=True)
    assert result['table_query']['value']==1
    assert 'scanned_rows' in captured[-1]['content']
    assert _unsupported_numbers('36.5 out of 50 marks. [S1]',[{'label':'S1','text':'Total (50M): 36.5'}])==[]


def test_removing_unindexed_upload_prevents_late_publication(service,tmp_path):
    path=workbook(tmp_path)
    pending=service.receive_import(path,None,path.name)
    original=service.embedder.encode
    def delete(*args,**kwargs):
        service.remove_document(pending['document_id'])
        return original(*args,**kwargs)
    service.embedder.encode=delete
    with pytest.raises(Exception):service.import_file(path,pending['document_id'],path.name)
    assert not service.documents()


def test_machine_shorthand_is_repaired_without_changing_the_score(service,tmp_path):
    from rag.answer import answer
    from dataclasses import replace
    doc=service.import_file(workbook(tmp_path))
    source=service.store.active_chunks([doc['document_id']])[0][0]
    source=replace(source,query_result={'columns':['Total (50M)']},text='Total (50M): 36.5')
    calls=[]
    def complete(messages,max_tokens):
        calls.append(messages)
        return {'result':{'status':'answered','answer':'36.5 out of 50M. [S1]' if len(calls)==1 else '36.5 out of 50 marks. [S1]'}}
    service.model.complete=complete
    result=answer('What is the mark?', [source],service.model)
    assert result['answer']=='36.5 out of 50 marks. [S1]'
    assert result['checks']['repair']['attempted'] and len(calls)==2


def test_upload_pending_does_not_break_retrieval_of_indexed_documents(service,tmp_path):
    from rag.retrieve import retrieve
    first=tmp_path/'notes.txt';first.write_text('Verified measurement 8.2 mm/s')
    service.import_file(first)
    service.receive_import(workbook(tmp_path),None,'scores.xlsx')
    assert retrieve(service.store,service.embedder,'measurement')


def test_numeric_aggregates_include_csv_numeric_text(tmp_path):
    from rag.tables import load_tables,execute_query
    path=tmp_path/'measurements.csv';path.write_text('name,value\nA,3.5\nB,6.5\nC,missing')
    result=execute_query(load_tables(path)[0],{'operation':'sum','columns':['value'],'filters':[]})
    assert result['value']==10 and result['numeric_rows']==2 and result['scanned_rows']==3


def test_query_identifier_rendering_preserves_exact_source_spelling():
    from rag.answer import _normalize_table_identifiers
    sources=[{'query_result':{'filters':[{'value':'24ALR001'}]}}]
    result=_normalize_table_identifiers({'answer':'24alr00_1 scored 36.5. 24ALR002 is another ID. [S1]'},sources)
    assert result['answer']=='24ALR001 scored 36.5. 24ALR002 is another ID. [S1]'
