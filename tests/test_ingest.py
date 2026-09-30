import hashlib
from types import SimpleNamespace
import pytest
from pypdf import PdfWriter
from backend.contracts import WorkbenchError, Extraction, Page
from rag.ingest import extract, chunk_pages

class Characters:
    def encode(self, text, add_special_tokens=False):
        return SimpleNamespace(offsets=[(i,i+1) for i in range(len(text))])

def test_text_bom_lines(tmp_path):
    path=tmp_path/'manual.txt'
    path.write_text('Title\nP-101: 7.1 mm/s.\n',encoding='utf-8-sig')
    result=extract(path)
    chunks=chunk_pages(result,Characters(),'d',path.name)
    assert chunks[0].line_start == 1 and chunks[0].line_end == 2
    assert not chunks[0].text.startswith('\ufeff')


def test_docx_paragraphs_tables_and_math_are_extracted(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    path=tmp_path/'problems.docx'
    doc=Document()
    doc.add_paragraph('Unit 3 problems')
    table=doc.add_table(rows=1,cols=2)
    table.cell(0,0).text='Pump A-17'
    table.cell(0,1).text='9.4 mm/s'
    paragraph=doc.add_paragraph()
    math=OxmlElement('m:oMath')
    value=OxmlElement('m:t'); value.text='x²+1'
    math.append(value); paragraph._p.append(math)
    doc.save(path)
    result=extract(path)
    text=result.pages[0].text
    assert 'Unit 3 problems' in text
    assert 'Pump A-17' in text and '9.4 mm/s' in text
    assert 'x²+1' in text
    assert result.pages[0].method=='docx_text'


def test_corrupt_docx_is_rejected(tmp_path):
    path=tmp_path/'bad.docx'; path.write_bytes(b'not a zip')
    with pytest.raises(WorkbenchError) as error: extract(path)
    assert error.value.code=='parse_failed'

def test_unicode_and_final_chunk():
    text='🌍测试abc'*100
    result=Extraction([Page(text,None)],[],hashlib.sha256(text.encode()).hexdigest())
    chunks=chunk_pages(result,Characters(),'d','a')
    assert len(chunks)==2
    assert chunks[-1].text.endswith('abc')
    assert all(c.text in text and len(c.text)<=384 for c in chunks)

@pytest.mark.parametrize('name,data,code',[('a.txt',b'','no_extractable_text'),('a.txt',b'\xff','invalid_encoding'),('a.exe',b'abc','unsupported_file'),('a.pdf',b'garbage','parse_failed')])
def test_bad_files(tmp_path,name,data,code):
    path=tmp_path/name; path.write_bytes(data)
    with pytest.raises(WorkbenchError) as error: extract(path)
    assert error.value.code==code

def test_limits_and_scanned_pdf(tmp_path):
    path=tmp_path/'blank.pdf'; writer=PdfWriter(); writer.add_blank_page(100,100); writer.write(path)
    with pytest.raises(WorkbenchError) as error: extract(path)
    assert error.value.code=='no_extractable_text'
    with pytest.raises(WorkbenchError) as error: extract(path,max_bytes=2)
    assert error.value.code=='file_too_large'
    with pytest.raises(WorkbenchError) as error: extract(path,max_pages=0)
    assert error.value.code=='page_limit'

def test_encrypted(tmp_path):
    path=tmp_path/'a.pdf'; w=PdfWriter(); w.add_blank_page(100,100); w.encrypt('secret'); w.write(path)
    with pytest.raises(WorkbenchError) as error: extract(path)
    assert error.value.code=='encrypted_pdf'

def test_page_boundaries():
    result=Extraction([Page('first',1),Page('second',3)],['page 2 needs OCR'],'v')
    chunks=chunk_pages(result,Characters(),'d','a.pdf')
    assert [c.page for c in chunks]==[1,3]


def test_office_spreadsheet_and_slide_text(tmp_path):
    from openpyxl import Workbook
    from pptx import Presentation
    book=Workbook();book.active.title='Measurements';book.active.append(['Pump',42])
    sheet=tmp_path/'measurements.xlsx';book.save(sheet)
    result=extract(sheet)
    assert result.pages[0].method=='spreadsheet_cells'
    assert 'B1: 42' in result.pages[0].text and 'Measurements' in result.pages[0].text
    assert 'saved values' in result.warnings[0]
    deck=Presentation();slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text='Inspection findings'
    slides=tmp_path/'inspection.pptx';deck.save(slides)
    result=extract(slides)
    assert result.pages[0].page==1 and result.pages[0].method=='slide_text'
    assert 'Inspection findings' in result.pages[0].text


def test_spreadsheet_ignores_inflated_dimensions_and_sparse_coordinates(tmp_path):
    from openpyxl import Workbook
    from zipfile import ZipFile, ZIP_DEFLATED
    from io import BytesIO
    book = Workbook()
    book.active['A1'] = 'First'
    book.active['XFD1048576'] = 'Last'
    buffer = BytesIO(); book.save(buffer)
    path = tmp_path / 'sparse.xlsx'
    with ZipFile(buffer) as original, ZipFile(path, 'w', ZIP_DEFLATED) as modified:
        for item in original.infolist():
            modified.writestr(item, original.read(item.filename))
    result = extract(path)
    assert 'A1: First' in result.pages[0].text
    assert 'XFD1048576: Last' in result.pages[0].text
    assert len(result.pages[0].text) < 100


@pytest.mark.parametrize('suffix', ['.tsv', '.yaml', '.xml', '.html', '.py', '.java', '.jsonl'])
def test_text_formats_and_utf16(tmp_path, suffix):
    path = tmp_path / ('source' + suffix)
    path.write_text('Earth 🌍\nSecond line', encoding='utf-16')
    assert extract(path).pages[0].text == 'Earth 🌍\nSecond line'


def test_binary_disguised_as_text_rejected(tmp_path):
    path = tmp_path / 'binary.txt'; path.write_bytes(b'abc\x00def')
    with pytest.raises(WorkbenchError, match='binary data'):
        extract(path)


def test_spreadsheet_shared_strings_and_invalid_references(tmp_path):
    import xlsxwriter
    from zipfile import ZipFile, ZIP_DEFLATED
    from io import BytesIO
    buffer=BytesIO()
    book=xlsxwriter.Workbook(buffer, {'in_memory': True})
    sheet=book.add_worksheet('Names'); sheet.write('A1','Earth'); book.close()
    path=tmp_path/'shared.xlsx'; path.write_bytes(buffer.getvalue())
    assert 'A1: Earth' in extract(path).pages[0].text
    with ZipFile(buffer) as original, ZipFile(path, 'w', ZIP_DEFLATED) as modified:
        for item in original.infolist():
            data=original.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml': data=data.replace(b'<v>0</v>',b'<v>-1</v>')
            modified.writestr(item,data)
    with pytest.raises(WorkbenchError, match='shared-string'):
        extract(path)
