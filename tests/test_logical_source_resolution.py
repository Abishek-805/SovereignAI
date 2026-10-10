import hashlib
from types import SimpleNamespace
import pytest
from rag.tables import _table,execute_query
from rag.identifiers import compile_exact_lookup
from rag.analytics import compile_measurements
from rag.table_resolution import logical_tables
from backend.contracts import WorkbenchError,Chunk
from backend.service import Workbench


def available():
    doc={'document_id':'workbook','active_hash':'v1','display_name':'Synthetic Professional Skills Training.xlsx'}
    return {'T15':(doc,_table('All',[['Member Id','WAT 1','WAT 2','CAT 1'],['24ALR056',7,6,12.5],['24ALR001',14,20,30]])),
            'T16':(doc,_table('Week 2',[['Member Id','Total'],['24ALR056',6],['24ALR001',20]]))}


def test_different_identifier_uses_same_summary_without_title_magic():
    for query,expected in [('24ALR056 professional skill training test results?',7),('Can you tell me the test results of 24ALR001?',14)]:
        source=logical_tables(available());plan=compile_exact_lookup(query,source)
        result=execute_query(source[plan['table']][1],plan)
        assert plan['table']=='T15' and result['records'][0]['WAT 1']==expected
        assert result['cell_locators'][0]['cells']['CAT 1'] in {'D2','D3'}
        assert plan['logical_resolution']['sheet']=='All'


def test_duplicate_indexes_and_overlapping_ranges_preserve_physical_rows():
    source=available();source['T17']=source['T15']
    table=source['T15'][1]
    source['T18']=(source['T15'][0],{**table,'records':[table['records'][1]],'record_rows':[3]})
    resolved=logical_tables(source)
    assert len(resolved)==2 and len(resolved['T15'][1]['records'])==2
    assert resolved['T15'][1]['internal_references']==['T15','T17','T18']
    result=execute_query(resolved['T15'][1],compile_exact_lookup('24ALR001 test results',resolved))
    assert result['matched_rows']==1 and result['record_locators']==[{'sheet':'All','row':3}]


def test_conflicting_same_row_is_not_deduplicated():
    source=available();doc,table=source['T15']
    source['T17']=(doc,{**table,'records':[{**table['records'][0],'WAT 1':999}],'record_rows':[2]})
    with pytest.raises(WorkbenchError,match='Conflicting source values'):logical_tables(source)


def test_distinct_sheets_do_not_become_interchangeable():
    source=available();doc=source['T15'][0]
    source['T17']=(doc,_table('Revision draft',[['Member Id','WAT 1'],['24ALR001',999]]))
    with pytest.raises(WorkbenchError,match='Conflicting WAT 1'):compile_exact_lookup('24ALR001 test results',logical_tables(source))
    result=compile_exact_lookup('24ALR001 Total results',logical_tables(available()))
    assert result['table']=='T16'


def test_genuinely_different_sources_and_versions_are_ambiguous():
    for change in [{'document_id':'other'},{'active_hash':'v2'}]:
        source=available();doc,table=source['T15'];source['T17']=({**doc,**change},table)
        with pytest.raises(WorkbenchError,match='Which source'):compile_exact_lookup('24ALR001 test results',logical_tables(source))


def test_missing_entity_and_assessment_remain_distinct():
    source=logical_tables(available())
    plan=compile_exact_lookup('24ALR999 test results',source)
    result=execute_query(source[plan['table']][1],plan)
    assert result['missing_entities']==['24ALR999'] and result['records']==[]
    with pytest.raises(WorkbenchError,match='assessment is unavailable'):compile_exact_lookup('24ALR001 WAT9 results',source)


def test_long_assessments_preserve_distinct_records():
    source={'T1':(available()['T15'][0],_table('Measurements',[['Key','Assessment','Score'],['24ALR001','Alpha',14],['24ALR001','Beta',20]]))}
    plan=compile_exact_lookup('24ALR001 test results',source)
    result=execute_query(source['T1'][1],plan)
    assert result['matched_rows']==2 and [r['Score'] for r in result['records']]==[14,20]
    assert [r['Assessment'] for r in result['records']]==['Alpha','Beta']


def test_distinct_measurement_sets_are_not_silently_omitted():
    source=available();source['T17']=(source['T15'][0],_table('Communication',[['Member Id','Communication'],['24ALR001',99]]))
    with pytest.raises(WorkbenchError,match='different measurement sets'):compile_exact_lookup('24ALR001 test results',source)


@pytest.fixture
def wb(tmp_path):
    from openpyxl import Workbook
    book=Workbook();book.active.title='All'
    for row in [['Member Id','WAT 1','WAT 2','CAT 1'],['24ALR056',7,6,12.5],['24ALR001',14,20,30]]:book.active.append(row)
    sheet=book.create_sheet('Week 2')
    for row in [['Member Id','Total'],['24ALR056',6],['24ALR001',20]]:sheet.append(row)
    path=tmp_path/'skills.xlsx';book.save(path)
    other=tmp_path/'other.csv';other.write_text('Member Id,WAT 1\n24ALR001,99\n')
    docs=[{'document_id':key,'active_hash':hashlib.sha256(file.read_bytes()).hexdigest(),'display_name':file.name,'source_extension':file.suffix,'chunk_count':1} for key,file in [('skills',path),('other',other)]]
    chunks={d['document_id']:Chunk(d['document_id'],d['document_id'],d['active_hash'],d['display_name'],'schema',None,1,2,retrieval_kind='schema') for d in docs}
    work=Workbench.__new__(Workbench);work.settings=SimpleNamespace(data_dir=tmp_path,context=16384,output_tokens=1024,safety_tokens=64)
    work.documents=lambda:docs;work.document_original=lambda key:path if key=='skills' else other
    work.store=SimpleNamespace(documents=lambda:docs,current=lambda key:next(d for d in docs if d['document_id']==key),
        get_chunk=lambda key:chunks[key],active_chunks=lambda scope:[(c,None) for k,c in chunks.items() if scope is None or k in scope])
    def forbidden(*a,**k):raise AssertionError('Model loop used for structured lookup')
    work.model=SimpleNamespace(complete=forbidden,count_messages=forbidden)
    return work


def test_verified_source_reuse_for_new_id_and_explicit_source_precedence(wb):
    first=wb._exact_document_lookup('24ALR056 test results in skills.xlsx',None)
    history=['User: '+first['original_question']]
    second=wb._exact_document_lookup('Can you tell me the test results of 24ALR001?',None,history,evidence_ids=['skills'])
    assert second['table_query']['records'][0]['WAT 1']==14 and second['operation_contract']['document_id']=='skills'
    assert second['table_query']['logical_resolution']['selection_reason']=='current_verified_prior_logical_table'
    average=wb._exact_document_lookup('tell me his overall avg',None,['User: '+second['original_question']],evidence_ids=['skills'])
    ranked=wb._exact_document_lookup('what are his best five marks?',None,['User: '+average['original_question']],evidence_ids=['skills'])
    assert [cell['value'] for cell in ranked['table_query']['measurements'][0]['selected']]==[30,20,14]
    explicit=wb._exact_document_lookup('24ALR001 test results in other.csv',None,history,evidence_ids=['skills'])
    assert explicit['table_query']['records'][0]['WAT 1']=='99' and explicit['operation_contract']['document_id']=='other'


def test_context_cannot_expand_permissions_or_reuse_stale_version(wb):
    first=wb._exact_document_lookup('24ALR056 test results in skills.xlsx',None)
    with pytest.raises(WorkbenchError,match='verified'):
        wb._exact_document_lookup('his overall avg',['other'],['User: '+first['original_question']],evidence_ids=['skills'])
    wb.store.documents=lambda:[{**d,'active_hash':'new'} for d in wb.documents()]
    with pytest.raises(WorkbenchError,match='verified'):
        wb._exact_document_lookup('his overall avg',None,['User: '+first['original_question']],evidence_ids=['skills'])


def test_unrelated_topic_does_not_pin_previous_workbook(wb):
    first=wb._exact_document_lookup('24ALR056 test results in skills.xlsx',None)
    with pytest.raises(WorkbenchError,match='Which source'):
        wb._exact_document_lookup('24ALR001 financial test results',None,['User: '+first['original_question']],evidence_ids=['skills'])


def test_prior_table_does_not_steal_explicit_assessment_field():
    source=logical_tables(available());prior={'document_id':'workbook','table':'All'}
    assert compile_exact_lookup('24ALR001 Total results',source,prior=prior)['table']=='T16'
    plan=compile_measurements('24ALR001 overall avg',source)
    assert execute_query(source[plan['table']][1],plan)['measurements'][0]['value']==64/3
