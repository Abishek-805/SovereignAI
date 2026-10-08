import json
from pathlib import Path
from backend.contracts import WorkbenchError
from backend.structured_output import normalize_json
from router.sandbox import safe_relative_name
from router.tool_registry import explicit_operation_requested
from rag.tables import compile_outcome_query,execute_query
from workflows.generated_documents import generate,validate_source_contract
import pytest

@pytest.mark.parametrize('name',['தமிழ்.txt','Tamil தமிழ்.txt','emoji😀.txt','safe name.txt'])
def test_unicode_path(name):assert safe_relative_name(name)
@pytest.mark.parametrize('name',['../bad','C:/bad','/bad','a\\bad','CON.txt','bad\x00.txt','a/b:bad','x\u202e.txt'])
def test_path_boundary(name):assert not safe_relative_name(name)
@pytest.mark.parametrize('goal',['Create a Flask project with auth.','Create a Node HTTP service with /health.','Build a Python CLI program.'])
def test_software_creation(goal):assert explicit_operation_requested(goal,'project_create')
@pytest.mark.parametrize('goal',['Run arbitrary PowerShell on the host.','Explain how to create a Flask project.','Do not create a Node HTTP service.','Create files in the current project.'])
def test_creation_boundary(goal):assert not explicit_operation_requested(goal,'project_create')
@pytest.mark.parametrize('text',['{}','[]','{"a":[1,null,"a\\\"b"]}','```json\n{"answer":42}\n```'])
def test_strict_json(text):json.loads(normalize_json(text))
@pytest.mark.parametrize('text',['Here is JSON: {}','{} and {}','{"a":1,"a":2}','{"a":NaN}','```json\n{}\n``` extra'])
def test_ambiguous_json(text):
 with pytest.raises(WorkbenchError):normalize_json(text)

def fixture_table():
 return {'name':'Results','columns':['Group','Score','Result'],'records':[{'Group':'A','Score':30,'Result':'FAIL'},{'Group':'A','Score':26,'Result':'FAIL'},{'Group':'A','Score':10,'Result':'PASS'},{'Group':'A','Score':None,'Result':'ABSENT'},{'Group':'A','Score':'bad','Result':'PASS'},{'Group':'A','Score':40,'Result':'PASS'}]}

def test_literal_numeric_and_combined_criteria():
 table=fixture_table();available={'T1':({'document_id':'doc'},table)};intent={'operation':'count','outcome':'pass','result_values':['PASS'],'entity_values':['A'],'entity_column':'Group','threshold':None,'score_columns':[],'scope':'all'}
 for request,expected in [('Count A records with Score at least 25.',3),('Count A records with Result PASS and Score at least 25.',1)]:
  plan=compile_outcome_query(intent,available,request);result=execute_query(table,plan);assert result['value']==expected;assert result['total']==6

def test_compound_aggregation():
 table=fixture_table();result=execute_query(table,{'operation':'aggregate','columns':[],'filters':[],'group_by':['Group'],'metrics':[{'operation':'count','column':None},{'operation':'sum','column':'Score'}]})
 assert [m['value'] for m in result['metrics']]==[6,106];assert result['groups'][0]['keys']=={'Group':'A'}

def test_xlsx_numeric_and_contract(tmp_path):
 from openpyxl import load_workbook
 p=tmp_path/'numbers.xlsx';generate(p,{'sheets':[{'name':'Counts','rows':[['Group','Count'],['A','7'],['B','0'],['C','100'],['D',None],['E','=SUM(B2:B4)']]}]});book=load_workbook(p,data_only=False);assert book.active['B2'].data_type=='n';assert book.active['B3'].value==0;assert book.active['B5'].value is None
 assert load_workbook(p,data_only=True).active['B6'].value==107
 source={'query_result':{'groups':[{'keys':{'Group':'A'},'metrics':[{'operation':'count','value':7}]}]}}
 assert validate_source_contract(p,[source])['source_contract_checked']
 q=tmp_path/'empty.xlsx';generate(q,{'sheets':[{'name':'Summary','rows':[['Group','Count']]}]})
 with pytest.raises(WorkbenchError):validate_source_contract(q,[source])
 assert not q.exists()


def test_source_contract_rejects_wrong_column_association(tmp_path):
 source={'query_result':{'groups':[{'keys':{'Group':'A'},'metrics':[{'operation':'count','value':7}]}]}}
 p=tmp_path/'wrong.xlsx';generate(p,{'sheets':[{'name':'Summary','rows':[['Count','Group'],['A',7]]}]})
 with pytest.raises(WorkbenchError):validate_source_contract(p,[source])

def test_compound_aggregate_request_uses_source_fields():
 from rag.tables import compile_aggregate_request
 table=fixture_table();available={'T1':({},table)}
 plan=compile_aggregate_request(available,'Summarize the number of records, cohort counts, and numeric Score totals.')
 assert plan['metrics']==[{'operation':'count','column':None},{'operation':'sum','column':'Score'}]
 assert plan['group_by']==['Group']
 assert compile_aggregate_request({'T1':({},table),'T2':({},table)},'Count records per Group') is None

def test_explicit_json_contract_scope():
 from backend.structured_output import apply_explicit_json_contract
 p={'messages':[{'role':'user','content':'Return exactly JSON with a nested array.'}]}
 assert apply_explicit_json_contract(p)['response_format']['type']=='json_object'
 assert 'response_format' not in apply_explicit_json_contract({'messages':[{'role':'user','content':'Explain JSON.'}]})


def test_draft_execution_snapshot_is_immutable_and_unpublished(tmp_path):
 from workflows.coding_workspace import CodingWorkspace
 from uuid import uuid4
 c=CodingWorkspace(tmp_path);w=c.create('Draft isolation')['workspace_id'];c.write(w,'notes.txt','KEEP')
 tid=uuid4().hex;r={'task_id':tid,'workspace_id':w,'state':'completed','checks':{},'changes':[{'path':'main.py','action':'create','before':None,'after':'print(42)'}]}
 c._stage_changes(w,r,c.raw_files(w))
 assert c.staged_files(w,tid)=={'notes.txt':b'KEEP','main.py':b'print(42)'}
 assert c.raw_files(w)=={'notes.txt':b'KEEP'}
 staged=c._directory(w)/'tasks'/tid/'staged'/'main.py';staged.write_text('print(0)')
 with pytest.raises(WorkbenchError,match='Draft bytes changed'):c.staged_files(w,tid)

def test_filename_placeholder_and_malformed_manifest_rejected():
 from workflows.source_quality import source_issue
 assert source_issue('main.py','main.py')
 assert source_issue('package.json','package.json')
 assert source_issue('package.json','not JSON')
 assert source_issue('package.json','{"main":"main.js"}') is None


def test_json_source_uses_native_object_not_double_serialization():
 import httpx
 from backend.model import LocalModel
 seen=[]
 def handler(request):
  if request.url.path=="/props":return httpx.Response(200,json={"default_generation_settings":{"n_ctx":8192}})
  if request.url.path=="/apply-template":return httpx.Response(200,json={"prompt":"formatted"})
  if request.url.path=="/tokenize":return httpx.Response(200,json={"tokens":[1]*100})
  payload=json.loads(request.content);seen.append(payload)
  return httpx.Response(200,json={"choices":[{"finish_reason":"stop","message":{"content":'{"main":"main.js","private":true}'}}]})
 model=LocalModel(transport=httpx.MockTransport(handler))
 try:
  result=model.complete_code([{"role":"user","content":json.dumps({"task":"Create a Node project","path":"package.json","language":"JSON"})}])
  assert json.loads(result["code"])=={"main":"main.js","private":True}
  assert seen[0]["response_format"]=={"type":"json_object"}
  assert "grammar" not in seen[0]
 finally:model.close()


def test_start_authorizes_execution_only_for_program_target():
 from router.tool_registry import explicit_operation_requested
 assert explicit_operation_requested("Start it and validate HTTP 200.","file_run")
 assert explicit_operation_requested("Launch the application.","file_run")
 assert not explicit_operation_requested("Start writing a document about Python.","file_run")
 assert not explicit_operation_requested("Do not start the server.","file_run")


def test_preserve_explicit_files_in_mixed_symbol_clause():
 from workflows.coding_workspace import CodingWorkspace
 names={"division.py":"", "notes.txt":"", "test_oracle.py":""}
 protected=CodingWorkspace._preserved_paths("Repair division.py. Preserve invoice_total(items), notes.txt, and test_oracle.py exactly.",names)
 assert protected=={"notes.txt","test_oracle.py"}

def test_project_plan_removes_only_redundant_workspace_root():
 from workflows.coding_workspace import CodingWorkspace
 ops=[{'action':'mkdir','path':'FullQA_Python_123'},
      {'action':'create','path':'FullQA_Python_123/main.py'},
      {'action':'create','path':'package/helpers.py'},
      {'action':'create','path':'../main.py'}]
 result=CodingWorkspace._relative_project_plan(ops,'Full QA Python 123')
 assert [item['path'] for item in result]==['main.py','package/helpers.py','../main.py']
 assert ops[1]['path']=='FullQA_Python_123/main.py'
