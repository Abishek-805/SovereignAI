import hashlib
import json
from pathlib import Path
from types import SimpleNamespace
import pytest
from docx import Document
from backend.contracts import WorkbenchError
from workflows.document_report import publish_report


@pytest.fixture
def settings(tmp_path):
    return SimpleNamespace(data_dir=tmp_path/'data')


@pytest.fixture
def answer():
    return {'status':'answered','answer':'The inspection records a completed review [S1].',
            'sources':[{'label':'S1','display_name':'Inspection.txt','page':1,'text':'The review was completed.'}]}


def locked(code=32):
    error=PermissionError('Injected Windows publication lock')
    error.winerror=code
    return error


def test_real_docx_and_manifest_are_published_together(settings,answer):
    result=publish_report(settings,'real-report','Summarize the inspection',answer)
    target=settings.data_dir.parent/'outputs'/'real-report'
    doc=Document(target/'Document report.docx')
    assert 'The inspection records a completed review [S1].' in [p.text for p in doc.paragraphs]
    assert 'The review was completed.' in [p.text for p in doc.paragraphs]
    manifest=json.loads((target/'manifest.json').read_text())
    assert manifest['files'][0]['sha256']==hashlib.sha256((target/'Document report.docx').read_bytes()).hexdigest()
    assert result=={'word':'/artifacts/real-report/word'}
    assert not list(target.parent.glob('.report-*'))


@pytest.mark.parametrize('code',[5,32,33])
def test_transient_windows_lock_retries_atomic_publication(settings,answer,monkeypatch,code):
    original=Path.rename;attempts=[];sleeps=[]
    target=settings.data_dir.parent/'outputs'/'retry-report'
    def rename(folder,destination):
        assert not target.exists()
        # Both completed files are present before the first publication attempt.
        assert Document(folder/'Document report.docx').paragraphs
        assert (folder/'manifest.json').is_file()
        attempts.append(destination)
        if len(attempts)<3:raise locked(code)
        return original(folder,destination)
    monkeypatch.setattr(Path,'rename',rename)
    monkeypatch.setattr('workflows.document_report.time.sleep',sleeps.append)
    publish_report(settings,'retry-report','Create a report',answer)
    assert len(attempts)==3 and sleeps==[0.05,0.1]
    assert Document(target/'Document report.docx').paragraphs


def test_permanent_lock_is_bounded_preserves_cause_and_cleans_staging(settings,answer,monkeypatch):
    error=locked();attempts=[];sleeps=[]
    def rename(folder,destination):attempts.append(destination);raise error
    monkeypatch.setattr(Path,'rename',rename)
    monkeypatch.setattr('workflows.document_report.time.sleep',sleeps.append)
    with pytest.raises(WorkbenchError,match='six attempts') as failure:
        publish_report(settings,'locked-report','Create a report',answer)
    assert failure.value.code=='artifact_publish' and failure.value.__cause__ is error
    assert len(attempts)==6 and sum(sleeps)==1.55
    outputs=settings.data_dir.parent/'outputs'
    assert not (outputs/'locked-report').exists() and not list(outputs.glob('.report-*'))


def test_non_windows_or_unrelated_io_error_is_not_retried(settings,answer,monkeypatch):
    error=OSError('No space left');sleeps=[];attempts=[]
    def rename(folder,destination):attempts.append(destination);raise error
    monkeypatch.setattr(Path,'rename',rename)
    monkeypatch.setattr('workflows.document_report.time.sleep',sleeps.append)
    with pytest.raises(WorkbenchError) as failure:publish_report(settings,'io-failure','Create a report',answer)
    assert failure.value.__cause__ is error and len(attempts)==1 and sleeps==[]


def test_existing_artifact_is_never_overwritten(settings,answer,monkeypatch):
    publish_report(settings,'existing','Original question',answer)
    target=settings.data_dir.parent/'outputs'/'existing'
    before={file.name:file.read_bytes() for file in target.iterdir()}
    def forbidden(*args):raise AssertionError('Existing destination must not be renamed over')
    monkeypatch.setattr(Path,'rename',forbidden)
    with pytest.raises(WorkbenchError,match='already exists') as failure:
        publish_report(settings,'existing','Replacement question',answer)
    assert failure.value.code=='artifact_publish'
    assert {file.name:file.read_bytes() for file in target.iterdir()}==before


def test_destination_appearing_between_retries_is_preserved(settings,answer,monkeypatch):
    target=settings.data_dir.parent/'outputs'/'concurrent';error=locked();attempts=[]
    def rename(folder,destination):
        attempts.append(destination);target.mkdir();(target/'marker.txt').write_text('other publication')
        raise error
    monkeypatch.setattr(Path,'rename',rename)
    monkeypatch.setattr('workflows.document_report.time.sleep',lambda _:None)
    with pytest.raises(WorkbenchError,match='already exists') as failure:
        publish_report(settings,'concurrent','Create a report',answer)
    assert len(attempts)==1 and failure.value.__cause__ is error
    assert (target/'marker.txt').read_text()=='other publication'
