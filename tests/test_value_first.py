import pytest
from rag.tables import bind_selection_request,execute_query,compile_outcome_query,_table
from rag.retrieve import document_scope_for_question
from backend.contracts import WorkbenchError
from rag.tables import resolve_query_table
from rag.tables import load_tables


@pytest.mark.parametrize('field',['Register Number','Unknown key','Third column'])
@pytest.mark.parametrize('operation',['select','count','sum','average','min','max'])
def test_wrong_identity_proposal_discovers_values(field,operation):
    table=_table('Professional Skills',[['Score','Other',field],[85,'x','24ALR056'],[78,'y','24ALR055']])
    plan=bind_selection_request({'operation':operation,'columns':['Score'],
        'filters':[{'column':'Member Id','operator':'eq','value':'24ALR056'}]},table,'Show 24ALR056 professional skill training results')
    result=execute_query(table,plan)
    assert result['matched_rows']==1
    assert result['record_locators']==[{'sheet':'Professional Skills','row':2}]
    assert plan['identity_column']==field
    assert result.get('value',85)==(1 if operation=='count' else 85)


def test_long_layout_and_legitimate_duplicate_identity():
    table=_table('Results',[['Person key','Assessment','Marks'],['24ALR056','Skills',85],['24ALR056','Communication',78]])
    plan=bind_selection_request({'operation':'select','columns':[],
        'filters':[{'column':'Member Id','operator':'eq','value':'24ALR056'},
                   {'column':'Assessment','operator':'eq','value':'Skills'}]},table,'Show 24ALR056 Skills')
    result=execute_query(table,plan)
    assert result['matched_rows']==1 and result['records'][0]['Marks']==85


@pytest.mark.parametrize('layout',['wide','long'])
def test_actual_workbook_reader_value_discovery(tmp_path,layout):
    from openpyxl import Workbook
    book=Workbook();sheet=book.active;sheet.title='Skills'
    rows=([['Register Number','Professional Skills','Communication'],['24ALR056',85,78]] if layout=='wide' else
          [['Register Number','Assessment','Marks'],['24ALR056','Professional Skills',85],['24ALR056','Communication',78]])
    for row in rows:sheet.append(row)
    path=tmp_path/'synthetic.xlsx';book.save(path);book.close()
    original=path.read_bytes()
    table=load_tables(path)[0]
    plan=bind_selection_request({'operation':'select','filters':[{'column':'Member Id','operator':'eq','value':'24ALR056'}]},table,'Show 24ALR056')
    result=execute_query(table,plan)
    assert result['matched_rows']==(1 if layout=='wide' else 2)
    assert result['record_locators'][0]=={'sheet':'Skills','row':2}
    assert path.read_bytes()==original


def test_unicode_case_and_whitespace_preserve_literal_digits():
    table=_table('Results',[['Key','Score'],[' 00Ab056 ',85]])
    result=execute_query(table,bind_selection_request({'operation':'select','filters':[]},table,'Show 00ab056'))
    assert result['matched_rows']==1 and result['records'][0]['Key']==' 00Ab056 '
    table=_table('Results',[['Key','Score'],['Café056',85]])
    result=execute_query(table,bind_selection_request({'operation':'select','filters':[]},table,'Show Cafe\u0301056'))
    assert result['matched_rows']==1 and result['missing_entities']==[]


def test_partial_identity_does_not_become_exact():
    table=_table('Results',[['Key','Score'],['24ALR056',85]])
    plan=bind_selection_request({'operation':'select','filters':[{'column':'Key','operator':'eq','value':'24ALR05'}]},table,'Show 24ALR05')
    assert execute_query(table,plan)['matched_rows']==0


def test_numeric_string_id_preserves_leading_zeros():
    table=_table('Results',[['Key','Score'],['000056',85],['56',10]])
    plan=bind_selection_request({'operation':'select','filters':[{'column':'Member Id','operator':'eq','value':'000056'}]},table,'Show 000056')
    result=execute_query(table,plan)
    assert result['matched_rows']==1 and result['records'][0]['Key']=='000056'


def test_conflicting_records_are_retained_with_separate_locators():
    table=_table('Results',[['Key','Score'],['24ALR056',85],['24ALR056',86]])
    plan=bind_selection_request({'operation':'select','filters':[]},table,'Show 24ALR056')
    result=execute_query(table,plan)
    assert [r['Score'] for r in result['records']]==[85,86]
    assert [r['row'] for r in result['record_locators']]==[2,3]


def test_multiple_ids_missing_sibling_and_ambiguous_numeric_preserved():
    table=_table('Results',[['Key','Score'],['24ALR056',85]])
    plan=bind_selection_request({'operation':'select','filters':[]},table,'24ALR056 and 55 professional skill training test results?')
    result=execute_query(table,plan)
    assert result['matched_rows']==1 and result['unresolved_identifiers']==['55']
    plan=bind_selection_request({'operation':'select','filters':[]},table,'Show 24ALR056 and 24ALR055')
    assert execute_query(table,plan)['missing_entities']==['24ALR055']


def test_ambiguous_identity_fields_do_not_guess():
    table=_table('Results',[['Key A','Key B','Score'],['24ALR056','24ALR056',85]])
    with pytest.raises(WorkbenchError,match='multiple source fields'):
        bind_selection_request({'operation':'select'},table,'Show 24ALR056')


def test_wrong_sheet_value_discovery_and_source_ambiguity():
    a=_table('Wrong sheet',[['Key','Score'],['24ALR001',10]])
    b=_table('Skills',[['Key','Score'],['24ALR056',85]])
    available={'T1':({'document_id':'a'},a),'T2':({'document_id':'a'},b)}
    plan={'table':'T1','operation':'select','columns':['Score']}
    assert resolve_query_table(plan,available,'Show 24ALR056')['table']=='T2'
    with pytest.raises(WorkbenchError):
        resolve_query_table(plan,{**available,'T3':({'document_id':'b'},b)},'Show 24ALR056')


def test_outcome_wrong_column_and_partial_rejection():
    table=_table('Skills',[['Register Number','Result'],['24ALR056','PASS']])
    intent={'operation':'count','scope':'all','result_values':['PASS'],'outcome':'pass',
        'entity_values':['24ALR056'],'entity_column':'Member Id'}
    plan=compile_outcome_query(intent,{'T1':({'document_id':'synthetic'},table)},'Count passes for 24ALR056')
    assert plan['filters'][0]['column']=='Register Number'
    assert execute_query(table,plan)['value']==1
    with pytest.raises(WorkbenchError):
        compile_outcome_query({**intent,'entity_values':['24ALR05']},{'T1':({'document_id':'synthetic'},table)},'Count passes for 24ALR05')


def test_assessment_name_cannot_replace_explicit_population():
    table=_table('Skills',[['Register Number','Assessment','Result'],['24ALR056','Professional Skills','PASS']])
    intent={'operation':'count','scope':'all','result_values':['PASS'],'outcome':'pass',
        'entity_values':['Professional Skills'],'entity_column':'Assessment'}
    with pytest.raises(WorkbenchError,match='source-bound user identifiers'):
        compile_outcome_query(intent,{'T1':({'document_id':'synthetic'},table)},'Count passes for 24ALR056 in Professional Skills')


def test_missing_measurement_and_unsupported_operation():
    table=_table('Results',[['Key','Score'],['24ALR056',None]])
    result=execute_query(table,bind_selection_request({'operation':'select','filters':[]},table,'Show 24ALR056'))
    assert result['records'][0]['Score'] is None
    with pytest.raises(WorkbenchError,match='Unsupported'):
        execute_query(table,{'operation':'execute_code'})


def test_explicit_file_scope_missing_and_duplicate():
    documents=[{'document_id':'a','display_name':'scores.csv'},{'document_id':'b','display_name':'other.csv'}]
    assert document_scope_for_question(documents,'Show results in scores.csv')==['a']
    with pytest.raises(WorkbenchError):document_scope_for_question(documents,'Show missing.csv')
    with pytest.raises(WorkbenchError):document_scope_for_question(documents+[{'document_id':'c','display_name':'scores.csv'}],'Show scores.csv')
    assert document_scope_for_question(documents,'Show scores.csv',['a'])==['a']
    assert set(document_scope_for_question(documents,'Compare scores.csv and other.csv'))=={'a','b'}
    assert document_scope_for_question([{'document_id':'c','display_name':'quarter results.csv'}],'Show quarter results.csv')==['c']
    assert document_scope_for_question([{'document_id':'c','display_name':'quarter results.csv'},{'document_id':'d','display_name':'results.csv'}],'Show quarter results.csv')==['c']
    with pytest.raises(WorkbenchError):document_scope_for_question(documents+[{'document_id':'c','display_name':'Scores.csv'}],'Show scores.csv')
    with pytest.raises(WorkbenchError):document_scope_for_question(documents,'Show notscores.csv')


def test_test_port_same_origin_keeps_foreign_origins_blocked():
    import asyncio
    from types import SimpleNamespace
    from backend.app import RequestBoundary
    async def probe(origin):
        sent=[]
        async def app(scope,receive,send):await send({'type':'http.response.start','status':204,'headers':[]})
        async def receive():return {'type':'http.request','body':b'','more_body':False}
        async def send(message):sent.append(message)
        boundary=RequestBoundary(app,SimpleNamespace(settings=SimpleNamespace(max_file_bytes=100)))
        await boundary({'type':'http','path':'/health','headers':[(b'host',b'127.0.0.1:8098'),(b'origin',origin.encode())]},receive,send)
        return sent[0]['status']
    assert asyncio.run(probe('http://127.0.0.1:8098'))==204
    assert asyncio.run(probe('http://127.0.0.1:9999'))==403
    assert asyncio.run(probe('https://evil.example'))==403
