import pytest
from backend.application_tools import ApplicationTools
from backend.contracts import WorkbenchError
from tests.test_service import service
from rag.tables import execute_query
from rag.tables import compile_outcome_query
from workflows.code_runtime import desktop_requirements


@pytest.mark.parametrize('goal',['Create an OpenCV project for face detection',
    'create a open cv project in python file the project is like with the laptop camera it should detect the users face is there are not'])
def test_program_project_keeps_selected_workspace(service,monkeypatch,goal):
    identifier=service.coding.create('Selected')['workspace_id']
    calls=[]
    monkeypatch.setattr(service,'run_coding_project_task',lambda workspace,*args,**kwargs:
        calls.append(workspace) or {'state':'completed'})
    operations=[{'tool':'project_create','target':'face_detection','value':'','input':''},
                {'tool':'file_edit','target':'main.py','value':'','input':''}]
    result=ApplicationTools(service,identifier,goal=goal).execute(operations,service.tasks.create('test',[]))
    assert result['workspace_id']==identifier and calls==[identifier]
    assert len(service.coding.list())==1


def test_explicit_separate_project_can_create_workspace(service):
    identifier=service.coding.create('Selected')['workspace_id']
    tool=ApplicationTools(service,identifier,goal='Create a separate project for face detection')
    result=tool._execute('project_create','New','','')
    assert result['workspace_id']!=identifier


def test_per_assessment_percentage_uses_cohort_and_independent_numerators():
    table={'name':'Tests','columns':['ID','Test 1','Test 2'],'records':[
        {'ID':'24ALR001','Test 1':'PASS','Test 2':'FAIL'},
        {'ID':'24ALR002','Test 1':'FAIL','Test 2':'PASS'},
        {'ID':'OTHER','Test 1':'PASS','Test 2':'PASS'}]}
    result=execute_query(table,{'operation':'percentage','columns':[],
        'filters':[{'column':'ID','operator':'contains','value':'24alr'}],
        'criteria':[{'column':column,'operator':'in','value':['PASS']} for column in ['Test 1','Test 2']]})
    assert [(s['passed'],s['total'],s['percentage']) for s in result['summaries']]==[(1,2,50.0),(1,2,50.0)]
    with pytest.raises(WorkbenchError,match='explicit result'):
        execute_query(table,{'operation':'percentage','columns':['Test 1'],'filters':[]})


def test_multi_column_numeric_aggregates_keep_separate_sample_counts():
    table={'name':'Tests','columns':['A','B'],'records':[{'A':10,'B':20},{'A':'AB','B':40}]}
    result=execute_query(table,{'operation':'average','columns':['A','B'],'filters':[]})
    assert result['summaries']==[{'column':'A','numeric_rows':1,'value':10.0},{'column':'B','numeric_rows':2,'value':30.0}]


def test_percentage_compiler_binds_prefix_across_complementary_assessments():
    doc={'document_id':'results'}
    available={
        'T1':(doc,{'name':'Summary','columns':['ID','Score'],'records':[{'ID':'TEAM01','Score':12}]}),
        'T2':(doc,{'name':'Exam A','columns':['ID','Result'],'records':[{'ID':'TEAM01','Result':'PASS'}]}),
        'T3':(doc,{'name':'Exam B','columns':['ID','Passed'],'records':[{'ID':'TEAM01','Passed':'FAIL'},{'ID':'TEAM02','Passed':'PASS'}]})}
    plan=compile_outcome_query({'result_values':['PASS','PASS'],'entity_values':['TEAM']},available,'Pass percentage of TEAM in each exam')
    assert plan['table']=='T2' and plan['filters']==[{'column':'ID','operator':'contains','value':'TEAM'}]
    assert [q['table'] for q in plan['additional_queries']]==['T3']
    assert plan['unavailable_assessments']==['Summary']
    count=compile_outcome_query({'operation':'count','result_values':['PASS'],'entity_values':['TEAM'],'scope':'one','assessments':['Exam B']},available,'How many TEAM students passed Exam B?')
    assert count['table']=='T3' and not count['additional_queries']
    assert execute_query(available['T3'][1],count)['value']==1
    with pytest.raises(WorkbenchError,match='current question'):
        compile_outcome_query({'result_values':['PASS'],'entity_values':['OTHER']},available,'Pass percentage of TEAM')


def test_desktop_capabilities_depend_on_calls_not_import_or_file_name():
    assert desktop_requirements('import cv2\ncv2.VideoCapture(0)\ncv2.imshow("Camera", frame)')==['camera','desktop window']
    assert desktop_requirements('import cv2\ncv2.VideoCapture("video.mp4")')==[]


def test_camera_run_returns_local_environment_command_before_docker(service,monkeypatch):
    identifier=service.coding.create('Camera')['workspace_id']
    service.coding.write(identifier,'main.py','import cv2\ncv2.VideoCapture(0)')
    monkeypatch.setattr(service,'_verified_coding_sandbox',lambda:pytest.fail('Desktop capability is checked before Docker availability'))
    with pytest.raises(WorkbenchError,match='run-workspace-python.ps1') as exc:
        service.execute_coding_file(identifier,'main.py')
    assert exc.value.code=='desktop_required'


def test_percentage_answer_keeps_denominator_and_unavailable_coverage(service,tmp_path):
    from openpyxl import Workbook
    book=Workbook(); summary=book.active;summary.title='Summary';summary.append(['ID','Score']);summary.append(['TEAM01',40])
    sheet=book.create_sheet('Exam');sheet.append(['ID','Result']);sheet.append(['TEAM01','PASS']);sheet.append(['TEAM02','FAIL'])
    path=tmp_path/'results.xlsx';book.save(path);doc=service.import_file(path)
    service.model.plan_table_query=lambda *args,**kwargs:{'operation':'percentage','_measure':{'result_values':['PASS'],'entity_values':['TEAM'],'scope':'all'}}
    service.model.complete=lambda *args,**kwargs:pytest.fail('Verified percentages do not need generated prose')
    result=service.ask('Pass percentage of TEAM in each exam',[doc['document_id']],force_documents=True)
    assert result['status']=='answered' and '50.0' in result['answer']
    assert 'denominator' in result['answer'] and 'Summary' in result['answer'] and '[S1]' in result['answer']


def test_outcome_polarity_prevents_mixed_pass_fail_numerator():
    table={'name':'Exam','columns':['ID','Result'],'records':[{'ID':'TEAM01','Result':'PASS'},{'ID':'TEAM02','Result':'FAIL'},{'ID':'TEAM03','Result':'AB'}]}
    plan=compile_outcome_query({'operation':'percentage','outcome':'pass','result_values':['PASS','FAIL','AB'],'entity_values':['TEAM']},{'T1':({'document_id':'d'},table)},'Pass percentage of TEAM')
    assert execute_query(table,plan)['summaries'][0]['passed']==1


def test_threshold_followup_preserves_cohort_and_excludes_absence_from_passes():
    table={'name':'All','columns':['ID','WAT 1','CAT 1'],'records':[
        {'ID':'TEAM01','WAT 1':25,'CAT 1':24.5},
        {'ID':'TEAM02','WAT 1':'AB','CAT 1':36.5},
        {'ID':'OTHER','WAT 1':40,'CAT 1':50}]}
    intent={'operation':'percentage','outcome':'pass','result_values':['PASS'],'entity_values':['TEAM'],'scope':'all','followup':True,'threshold':25,'threshold_operator':'gte','score_columns':['WAT 1','CAT 1']}
    plan=compile_outcome_query(intent,{'T1':({'document_id':'d'},table)},'Below 25 is failed; give the report for all tests',['user: pass percentage of TEAM in each test'])
    summaries=execute_query(table,plan)['summaries']
    assert [(r['passed'],r['total'],r['percentage']) for r in summaries]==[(1,2,50.0),(1,2,50.0)]
    with pytest.raises(WorkbenchError,match='supplied'):
        compile_outcome_query({**intent,'threshold':99},{'T1':({'document_id':'d'},table)},'Below 25 fails',['user: pass percentage of TEAM'])


def test_generic_report_does_not_select_report_named_document():
    from rag.retrieve import document_scope_for_question
    docs=[{'document_id':'workshop','display_name':'workbench-report-acceptance.txt'}, {'document_id':'marks','display_name':'PST Marks.xlsx'}]
    question='Students below 25 failed; give report accordingly for all the tests'
    assert document_scope_for_question(docs,question,['workshop','marks'])==['workshop','marks']
    assert document_scope_for_question(docs,'Use document PST Marks.xlsx',['workshop','marks'])==['marks']


def test_model_threshold_refinement_reuses_prior_measure_and_complements_rule(monkeypatch):
    import json
    from backend.model import LocalModel
    model=LocalModel()
    refinement={'operation':'select','outcome':'pass','followup':True,'threshold':25,'threshold_operator':'lt','rule_outcome':'fail','score_columns':['Score'],'result_values':['PASS'],'entity_values':[],'scope':'all','assessments':[]}
    previous={**refinement,'operation':'percentage','followup':False,'threshold':None,'threshold_operator':'none','rule_outcome':'other','entity_values':['TEAM']}
    replies=iter([previous,refinement])
    monkeypatch.setattr(model,'_request',lambda *args,**kwargs:{'choices':[{'finish_reason':'stop','message':{'content':json.dumps(next(replies))}}]})
    question='Below 25 fails; update the report'
    query=model.plan_table_query(question,[{'id':'T1','sheet':'Exam','columns':['ID','Score'],'categorical_values':{'Result':['PASS','FAIL']}}],['user: Pass percentage for TEAM in all tests','user: '+question])
    assert query['operation']=='percentage'
    assert query['_measure']['entity_values']==['TEAM']
    assert query['_measure']['threshold_operator']=='gte'


def test_literal_source_cohort_is_retained_when_model_omits_it():
    table={'name':'Exam','columns':['ID','Result'],'records':[{'ID':'24ALR001','Result':'PASS'},{'ID':'OTHER','Result':'PASS'}]}
    plan=compile_outcome_query({'operation':'percentage','outcome':'pass','entity_values':[]},{'T1':({'document_id':'d'},table)},'Pass percentage of all 24alr')
    assert execute_query(table,plan)['matched_rows']==1
