"""Existing evidence and deterministic artifact containers exposed as kernel tools."""
import hashlib
import json
import re
from pathlib import Path
from uuid import uuid4
from workflows.generated_documents import generate

class KernelCapabilities:
    def __init__(self,service,request,document_ids=None):
        self.service=service;self.request=request;self.document_ids=document_ids
        self.root=service.settings.data_dir.parent/'outputs'/('agent-artifacts-'+uuid4().hex)
        self.created={}
    def evidence(self,query):
        from dataclasses import asdict
        chunks=self.service.store.active_chunks(self.document_ids)
        chunks=[asdict(chunk) for chunk,_vector in chunks]
        words=set(re.findall(r'\w+',query.casefold()))
        ranked=sorted(chunks,key=lambda c:sum(w in json.dumps(c,default=str).casefold() for w in words),reverse=True)
        return {'evidence':ranked[:12],'scope':self.document_ids,'truncated':len(chunks)>12}
    def create(self,kind,encoded):
        suffix={'create_docx':'.docx','create_xlsx':'.xlsx','create_pptx':'.pptx'}[kind]
        if not re.search(r'\b(?:create|make|generate|export|save)\b',self.request,re.I) or not re.search(r'\b(?:document|report|word|excel|spreadsheet|slides|presentation|docx|xlsx|pptx)\b',self.request,re.I):
            raise ValueError('Artifact creation is not authorized by the original request')
        payload=json.loads(encoded);name=payload['filename']
        if Path(name).name!=name or Path(name).suffix!=suffix:raise ValueError('Expected a simple artifact filename with matching extension')
        self.root.mkdir(parents=True,exist_ok=True);path=self.root/name
        if kind=='create_pptx':
            from pptx import Presentation
            presentation=Presentation()
            for section in payload['specification']['sections'][:64]:
                slide=presentation.slides.add_slide(presentation.slide_layouts[1]);slide.shapes.title.text=str(section.get('heading',''))
                slide.placeholders[1].text='\n'.join(map(str,section.get('paragraphs',[])))
            presentation.save(path)
            if not Presentation(path).slides:raise ValueError('Empty presentation')
            validation={'opened':True,'slides':len(presentation.slides)}
        else:validation=generate(path,payload['specification'])
        digest=hashlib.sha256(path.read_bytes()).hexdigest();self.created[name]={'sha256':digest,'validation':validation}
        (self.root/'manifest.json').write_text(json.dumps({'files':self.created},indent=2))
        return {'path':str(path),'sha256':digest,'validation':validation,'container_only':True,'semantic_validation':'not inferred from container readback'}
    def validate(self,name):
        if name not in self.created:raise ValueError('Artifact is not owned by this session')
        actual=hashlib.sha256((self.root/name).read_bytes()).hexdigest()
        return {'unchanged':actual==self.created[name]['sha256'],**self.created[name]}
    def tools(self):
        return {'query_rag':self.evidence,'retrieve_evidence':self.evidence,
                **{kind:(lambda request,kind=kind:self.create(kind,request)) for kind in ['create_docx','create_xlsx','create_pptx']},
                'validate_artifact':self.validate}
