"""Export a grounded document answer as a reviewable Word report."""
import hashlib
import json
import re
import time
from tempfile import TemporaryDirectory
from pathlib import Path
from docx import Document
from backend.contracts import WorkbenchError
from router.telemetry import measured_validation


def _publish_directory(folder, target):
    """Publish once atomically; tolerate bounded Windows scanner/file-handle locks."""
    delays=(0.05,0.1,0.2,0.4,0.8)
    last_error=None
    for attempt in range(len(delays)+1):
        if target.exists():
            raise WorkbenchError('artifact_publish','Report destination already exists; no existing artifact was replaced.') from last_error
        try:
            # Windows directory rename is atomic and refuses an existing target.
            folder.rename(target)
            return
        except OSError as error:
            last_error=error
            if getattr(error,'winerror',None) in {5,32,33} and attempt<len(delays):
                time.sleep(delays[attempt])
                continue
            detail='Report publication failed'
            if getattr(error,'winerror',None) in {5,32,33}:
                detail+=' because Windows kept the report folder locked after six attempts'
            else:
                detail+=': '+str(error)
            raise WorkbenchError('artifact_publish',detail+'. No partial report was published.') from error


def is_overview_request(question):
    return bool(re.fullmatch(
        r'(?:please\s+)?(?:generate|create|write|make|prepare|draft)(?:\s+me)?\s+(?:a\s+)?(?:word\s+)?report(?:\s+(?:on|about|from|of)\s+(?:(?:the|these|those|my)\s+)?(?:selected\s+)?(?:documents|files))?[.! ]*',
        question.strip(), re.I))


def selected_document_overview(chunks, question):
    """Build a cited inventory when no narrower report topic was supplied."""
    first_by_document = {}
    for chunk, _vector in chunks:
        first_by_document.setdefault(chunk.document_id, chunk)
    if not first_by_document:
        raise WorkbenchError('missing_evidence', 'Select documents with readable text before creating a report.')
    sources = []
    paragraphs = [f'Overview of {len(first_by_document)} selected document(s). Each entry below is an excerpt from the indexed file; review the source for full context.']
    for chunk in sorted(first_by_document.values(), key=lambda item: item.display_name.casefold()):
        label = f'S{len(sources) + 1}'
        excerpt = re.sub(r'\s+', ' ', chunk.text).strip()[:600]
        if not excerpt:
            continue
        sources.append({**chunk.to_dict(), 'label': label})
        paragraphs.append(f'{chunk.display_name}: {excerpt}{"…" if len(chunk.text) > 600 else ""} [{label}]')
    if not sources:
        raise WorkbenchError('missing_evidence', 'The selected documents have no readable text for a report.')
    return {'status': 'answered', 'answer': '\n\n'.join(paragraphs), 'sources': sources,
            'checks': {'citation_ids_valid': True, 'semantic_support': 'source_excerpts'},
            'timings': {}, 'model': None,
            'routing': {'capability': 'document_report', 'model': 'No model used',
                        'reason': 'Cited excerpts from the selected files'}}


@measured_validation
def publish_report(settings, task_id, question, result):
    if result.get('status') != 'answered' or not result.get('sources'):
        detail=result.get('answer','')
        raise WorkbenchError('missing_evidence','The documents did not support a grounded report. Inspect the answer and sources first.'+(' '+detail[:1000] if isinstance(detail,str) and detail else ''))
    root = settings.data_dir.parent / 'outputs'
    root.mkdir(parents=True,exist_ok=True)
    with TemporaryDirectory(prefix='.report-',dir=root) as directory:
        folder = Path(directory)
        path = folder / 'Document report.docx'
        doc = Document()
        doc.add_heading('Document report',0)
        doc.add_paragraph(question)
        doc.add_paragraph('AI-assisted draft. Review the cited evidence before use.')
        for paragraph in result['answer'].split('\n'):
            if paragraph.strip(): doc.add_paragraph(paragraph)
        doc.add_heading('Sources',1)
        for source in result['sources']:
            doc.add_paragraph(f"[{source['label']}] {source['display_name']} · page {source.get('page') or 'not recorded'}")
            doc.add_paragraph(source['text'])
        coverage=result.get('coverage')
        if isinstance(coverage,dict) and coverage.get('mode')=='overview':
            doc.add_heading('Evidence coverage',1)
            documents=coverage.get('documents',[])
            represented=sum(bool(item['included_passages']) for item in documents)
            included=sum(item['included_passages'] for item in documents)
            indexed=sum(item['indexed_passages'] for item in documents)
            doc.add_paragraph(f'{represented} of {len(documents)} connected documents are represented by '
                              f'{included} of {indexed} indexed passages in this report.')
            doc.add_paragraph('This is a synthesis of indexed excerpts. Indexed passage coverage does not '
                              'verify that all original file content, images, tables, or layout were interpreted.')
            omitted=[item['display_name'] for item in documents if not item['included_passages']]
            if omitted:
                doc.add_paragraph('Documents without included evidence: '+ '; '.join(omitted))
            partial=[item for item in documents if 0<item['included_passages']<item['indexed_passages']]
            if partial:
                doc.add_paragraph('Documents represented by partial indexed evidence: '+ '; '.join(
                    f"{item['display_name']} ({item['included_passages']} of {item['indexed_passages']} passages)"
                    for item in partial))
        doc.save(path)
        if not Document(path).paragraphs:
            raise WorkbenchError('artifact_invalid','Report could not be read back')
        metadata={'task_id':task_id,'created_at':time.time(),'files':[{'name':path.name,
                  'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'bytes':path.stat().st_size}]}
        (folder/'manifest.json').write_text(json.dumps(metadata,indent=2),encoding='utf-8')
        _publish_directory(folder,root/task_id)
    return {'word':f'/artifacts/{task_id}/word'}
