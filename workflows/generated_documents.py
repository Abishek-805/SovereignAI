"""Deterministic document containers; model output is data, never executable code."""
import json,re
from pathlib import Path
from backend.contracts import WorkbenchError


def generate(path, specification):
    path=Path(path)
    if not isinstance(specification,dict):raise WorkbenchError('generation_format','Expected a structured document object')
    if path.exists():raise WorkbenchError('workspace_conflict','Generated document already exists')
    suffix=path.suffix.lower()
    try:
        if suffix=='.docx':
            from docx import Document
            doc=Document()
            title=specification.get('title','')
            if title:doc.add_heading(str(title),0)
            sections=specification.get('sections',[])
            if not isinstance(sections,list) or not sections:raise ValueError('At least one section required')
            for section in sections[:64]:
                if section.get('heading'):doc.add_heading(str(section['heading']),1)
                for paragraph in section.get('paragraphs',[])[:128]:doc.add_paragraph(str(paragraph))
                for rows in section.get('tables',[])[:16]:
                    if isinstance(rows,dict):
                        headers=rows.get('headers',rows.get('columns',[]))
                        rows=([headers] if headers else [])+rows.get('rows',[])
                    if not rows or not isinstance(rows,list) or any(not isinstance(row,list) for row in rows):raise ValueError('Invalid table')
                    width=max(len(row) for row in rows)
                    table=doc.add_table(rows=len(rows),cols=width)
                    table.style='Table Grid'
                    for i,row in enumerate(rows):
                        for j,value in enumerate(row):table.cell(i,j).text=str(value)
            doc.save(path)
            reopened=Document(path)
            if not any(p.text.strip() for p in reopened.paragraphs):raise ValueError('Document is empty')
            return {'opened':True,'paragraphs':len(reopened.paragraphs),'tables':len(reopened.tables)}
        if suffix=='.xlsx':
            from openpyxl import Workbook,load_workbook
            workbook=Workbook();workbook.remove(workbook.active)
            sheets=specification.get('sheets',[])
            if not isinstance(sheets,list) or not sheets:raise ValueError('At least one sheet required')
            for sheet in sheets[:32]:
                ws=workbook.create_sheet(str(sheet['name']))
                rows=sheet.get('rows',[])
                if not rows:raise ValueError('Sheet rows required')
                for row in rows[:10000]:
                    if not isinstance(row,list) or len(row)>256:raise ValueError('Invalid worksheet row')
                    for value in row:
                        if not isinstance(value,(str,int,float,bool,type(None))):raise ValueError('Unsupported cell type')
                        if isinstance(value,str) and value.startswith('=') and ('[' in value or '://' in value or 'DDE' in value.upper()):raise ValueError('External formulas are not supported')
                    # Preserve numeric data even when a model serializes a
                    # numeric column as strings. Identifiers and zero-padded
                    # codes remain text; blanks retain their missing value.
                    headers = rows[0]
                    normalized = []
                    for index, value in enumerate(row):
                        header = str(headers[index]).casefold() if index < len(headers) else ''
                        numeric = bool(re.search(r'\b(count|score|quantity|amount|total|price|rate|percent|balance)\b', header))
                        if row is not rows[0] and numeric and isinstance(value,str) and re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?',value.strip()):
                            value = float(value) if '.' in value else int(value)
                        if numeric and re.search(r'\b(count|quantity)\b',header) and isinstance(value,(int,float)) and not isinstance(value,bool) and value < 0:
                            raise ValueError('Count/quantity cells cannot be negative')
                        normalized.append(value)
                    ws.append(normalized)
                ws.freeze_panes='A2'
            calculations={}
            for ws in workbook:
                for row in ws:
                    for cell in row:
                        formula=cell.value
                        if not isinstance(formula,str) or not formula.startswith('='):continue
                        match=re.fullmatch(r'=SUM\(([A-Z]+[1-9][0-9]*):([A-Z]+[1-9][0-9]*)\)',formula,re.I)
                        if not match:raise ValueError('Unsupported formula in '+ws.title+'!'+cell.coordinate+': '+formula+'; use SUM over numeric cells')
                        values=[item.value for cells in ws[match.group(1)+':'+match.group(2)] for item in cells]
                        numbers=[value for value in values if isinstance(value,(int,float)) and not isinstance(value,bool)]
                        if not numbers or len(numbers)!=len(values):raise ValueError('Formula '+formula+' references empty or nonnumeric cells at '+ws.title+'!'+cell.coordinate+'. Actual numeric cells: '+', '.join(c.coordinate+'='+str(c.value) for r in ws for c in r if isinstance(c.value,(int,float)) and not isinstance(c.value,bool))+'. Repair the formula reference; preserve the requested rows and columns.')
                        calculations[ws.title+'!'+cell.coordinate]=sum(numbers)
            workbook.save(path)
            # Preserve formulas and deterministic calculated values for offline readers.
            if calculations:
                from zipfile import ZipFile,ZIP_DEFLATED
                import xml.etree.ElementTree as ET
                namespace='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
                with ZipFile(path) as archive:parts={name:archive.read(name) for name in archive.namelist()}
                for index,ws in enumerate(workbook,1):
                    name=f'xl/worksheets/sheet{index}.xml'
                    root=ET.fromstring(parts[name])
                    for cell in root.iter(namespace+'c'):
                        key=ws.title+'!'+cell.attrib['r']
                        if key not in calculations:continue
                        value=cell.find(namespace+'v')
                        if value is None:value=ET.SubElement(cell,namespace+'v')
                        value.text=str(calculations[key])
                    parts[name]=ET.tostring(root,encoding='utf-8')
                with ZipFile(path,'w',ZIP_DEFLATED) as archive:
                    for name,data in parts.items():archive.writestr(name,data)
            check=load_workbook(path,data_only=False)
            return {'opened':True,'sheets':check.sheetnames,'computed_formula_values':calculations,'formula_cache_written':True}
        raise ValueError('Supported generated containers are DOCX and XLSX')
    except (ValueError,KeyError,TypeError) as error:
        path.unlink(missing_ok=True)
        raise WorkbenchError('generation_format',str(error)) from error


def validate_source_contract(path, sources):
    """Verify authoritative grouped results, not merely a readable container."""
    if Path(path).suffix.lower()!='.xlsx':return {'source_contract_checked':False}
    from openpyxl import load_workbook
    book=load_workbook(path,data_only=True)
    required=[]
    for source in sources:
        query=source.get('query_result') or {}
        for group in query.get('groups',[]):
            for metric in group.get('metrics',[]):
                if metric.get('operation')=='count':
                    required.append([*group['keys'].values(),metric['value']])
    for source in sources:
        for group in (source.get('query_result') or {}).get('groups',[]):
            for metric in group.get('metrics',[]):
                if metric.get('operation')!='count':continue
                matched=False
                for sheet in book:
                    rows=list(sheet.iter_rows(values_only=True))
                    if not rows:continue
                    headers=[str(v).strip().casefold() for v in rows[0]]
                    keys=group['keys']
                    if any(k.casefold() not in headers for k in keys):continue
                    count_indices=[i for i,h in enumerate(headers) if h in {'count','record count','records','number of records'}]
                    for row in rows[1:]:
                        if all(row[headers.index(k.casefold())]==v for k,v in keys.items()) and any(isinstance(row[i],(int,float)) and not isinstance(row[i],bool) and row[i]==metric['value'] for i in count_indices):matched=True
                if not matched:
                    Path(path).unlink(missing_ok=True)
                    raise WorkbenchError('generation_format','Required source summary row missing or incorrectly typed/associated: '+json.dumps({'keys':group['keys'],'count':metric['value']}))
    return {'source_contract_checked':bool(required),'required_summary_rows':required}


def specification_schema(suffix):
    cell={'type':['string','number','boolean','null']}
    rows={'type':'array','minItems':1,'items':{'type':'array','items':cell}}
    if suffix=='.xlsx':
        item={'type':'object','properties':{'name':{'type':'string'},'rows':rows},'required':['name','rows'],'additionalProperties':False}
        return {'type':'object','properties':{'sheets':{'type':'array','minItems':1,'items':item}},'required':['sheets'],'additionalProperties':False}
    section={'type':'object','properties':{'heading':{'type':'string'},'paragraphs':{'type':'array','items':{'type':'string'}},'tables':{'type':'array','items':rows}},'required':['heading','paragraphs','tables'],'additionalProperties':False}
    return {'type':'object','properties':{'title':{'type':'string'},'sections':{'type':'array','minItems':1,'items':section}},'required':['title','sections'],'additionalProperties':False}
